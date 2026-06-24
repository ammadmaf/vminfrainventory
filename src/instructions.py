"""Instructions worksheet generation for Virtualization Administration Toolkit."""

from __future__ import annotations

from copy import copy
from dataclasses import dataclass
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from .styles import header_style, table_styles, workbook_theme


INSTRUCTION_COLUMNS: tuple[str, ...] = (
    "Step",
    "Area",
    "Instruction",
    "Owner",
)


@dataclass(frozen=True)
class InstructionRecord:
    """Workbook usage instruction row."""

    step: int
    area: str
    instruction: str
    owner: str

    def as_row(self) -> tuple[Any, ...]:
        """Return the record in worksheet column order."""
        return (self.step, self.area, self.instruction, self.owner)


class InstructionsWorksheet:
    """Builds the workbook instructions worksheet."""

    sheet_name = "Instructions"
    table_name = "WorkbookInstructionsTable"

    def __init__(self, records: list[InstructionRecord] | None = None) -> None:
        """Initialize instruction worksheet data and styles."""
        self.records = records or self._default_records()
        self.theme = workbook_theme()
        self.header = header_style("instructions_header")
        self.table = table_styles()

    def build(self, workbook: Workbook) -> Worksheet:
        """Create and format the Instructions worksheet."""
        worksheet = self._get_worksheet(workbook)
        self._reset_worksheet(worksheet)
        self._configure_worksheet(worksheet)
        self._add_title(worksheet)
        self._write_headers(worksheet)
        self._write_rows(worksheet)
        self._apply_table(worksheet)
        self._apply_body_formatting(worksheet)
        self._auto_fit_columns(worksheet)
        return worksheet

    def _get_worksheet(self, workbook: Workbook) -> Worksheet:
        """Return the Instructions worksheet."""
        if self.sheet_name in workbook.sheetnames:
            return workbook[self.sheet_name]
        return workbook.create_sheet(self.sheet_name)

    @staticmethod
    def _reset_worksheet(worksheet: Worksheet) -> None:
        """Clear existing content and merged ranges."""
        for merged_range in list(worksheet.merged_cells.ranges):
            worksheet.unmerge_cells(str(merged_range))
        if worksheet.max_row:
            worksheet.delete_rows(1, worksheet.max_row)
        worksheet.tables.clear()

    def _configure_worksheet(self, worksheet: Worksheet) -> None:
        """Apply worksheet layout settings."""
        worksheet.sheet_view.showGridLines = False
        worksheet.freeze_panes = "A5"
        worksheet.sheet_properties.tabColor = self.theme.palette.blue
        worksheet.row_dimensions[1].height = 34
        worksheet.row_dimensions[4].height = 28

    def _add_title(self, worksheet: Worksheet) -> None:
        """Add workbook instructions title."""
        worksheet.merge_cells("A1:D2")
        cell = worksheet["A1"]
        cell.value = "Virtualization Administration Toolkit - Instructions"
        cell.font = Font(
            name="Calibri Light",
            size=22,
            bold=True,
            color=self.theme.palette.white,
        )
        cell.fill = PatternFill("solid", fgColor=self.theme.palette.navy)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    def _write_headers(self, worksheet: Worksheet) -> None:
        """Write instruction table headers."""
        for column_index, header in enumerate(INSTRUCTION_COLUMNS, start=1):
            cell = worksheet.cell(row=4, column=column_index, value=header)
            self._apply_named_style(cell, self.header)

    def _write_rows(self, worksheet: Worksheet) -> None:
        """Write workbook usage instructions."""
        for row_index, record in enumerate(self.records, start=5):
            for column_index, value in enumerate(record.as_row(), start=1):
                worksheet.cell(row=row_index, column=column_index, value=value)

    def _apply_table(self, worksheet: Worksheet) -> None:
        """Create an Excel table over the instructions."""
        table_range = f"A4:D{worksheet.max_row}"
        table = Table(displayName=self.table_name, ref=table_range)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        worksheet.add_table(table)

    def _apply_body_formatting(self, worksheet: Worksheet) -> None:
        """Apply table styles and wrapping to instruction rows."""
        for row in worksheet.iter_rows(
            min_row=5,
            max_row=worksheet.max_row,
            min_col=1,
            max_col=len(INSTRUCTION_COLUMNS),
        ):
            for cell in row:
                style = (
                    self.table["table_numeric"]
                    if cell.column == 1
                    else self.table["table_body"]
                )
                self._apply_named_style(cell, style)
                if cell.column == 3:
                    cell.alignment = Alignment(
                        horizontal="left",
                        vertical="top",
                        wrap_text=True,
                    )
            worksheet.row_dimensions[row[0].row].height = 42

    @staticmethod
    def _apply_named_style(cell: Cell, named_style: Any) -> None:
        """Copy a named style's visual attributes onto a cell."""
        cell.font = copy(named_style.font)
        cell.fill = copy(named_style.fill)
        cell.border = copy(named_style.border)
        cell.alignment = copy(named_style.alignment)
        cell.number_format = named_style.number_format

    @staticmethod
    def _auto_fit_columns(worksheet: Worksheet) -> None:
        """Set readable column widths."""
        widths = {"A": 10, "B": 24, "C": 78, "D": 26}
        for column, width in widths.items():
            worksheet.column_dimensions[column].width = width

    @staticmethod
    def _default_records() -> list[InstructionRecord]:
        """Return buyer-facing workbook usage instructions."""
        return [
            InstructionRecord(
                1,
                "Workbook",
                "Review the Cover and Instructions sheets before updating data.",
                "Workbook Owner",
            ),
            InstructionRecord(
                2,
                "VM Inventory",
                "Replace sample VM records with exported Virtualization inventory data.",
                "Virtualization Administrator",
            ),
            InstructionRecord(
                3,
                "Operations",
                "Update host, datastore, backup, snapshot, license, and capacity "
                "worksheets on the same reporting cycle.",
                "Infrastructure Operations",
            ),
            InstructionRecord(
                4,
                "Dashboard",
                "Open the workbook in Excel and recalculate formulas to refresh "
                "dashboard KPIs and chart values.",
                "Report Consumer",
            ),
            InstructionRecord(
                5,
                "Commercial Use",
                "Customize branding, sample data, and support details before "
                "client delivery or marketplace publication.",
                "Product Owner",
            ),
        ]


