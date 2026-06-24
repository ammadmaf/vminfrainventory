"""Datastore capacity worksheet generation.

This module creates only the ``Datastore`` worksheet and includes Excel
formulas for free capacity, percentage used, and utilization warnings.
"""

from __future__ import annotations

from copy import copy
from dataclasses import dataclass
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from .styles import header_style, table_styles, workbook_theme


DATASTORE_COLUMNS: tuple[str, ...] = (
    "Datastore",
    "Type",
    "Cluster",
    "Capacity (GB)",
    "Used (GB)",
    "Free (GB)",
    "Percentage Used",
    "Warnings",
    "Notes",
    "Assigned VM Count",
    "Deletion Status",
)


@dataclass(frozen=True)
class DatastoreRecord:
    """Sample datastore capacity row."""

    datastore: str
    datastore_type: str
    cluster: str
    capacity_gb: int
    used_gb: int
    notes: str

    def as_row(self, row_index: int) -> tuple[Any, ...]:
        """Return the record in worksheet column order with formulas."""
        free_formula = f"=D{row_index}-E{row_index}"
        used_percent_formula = f"=IFERROR(E{row_index}/D{row_index},0)"
        warning_formula = (
            f'=IF(G{row_index}>=0.9,"Critical",'
            f'IF(G{row_index}>=0.8,"Warning","Healthy"))'
        )
        assigned_count_formula = (
            f"=COUNTIF('VM Inventory'!$J:$J,A{row_index})"
        )
        deletion_status_formula = (
            f'=IF(J{row_index}>0,'
            '"Assigned to VM - must not delete","Safe to delete")'
        )

        return (
            self.datastore,
            self.datastore_type,
            self.cluster,
            self.capacity_gb,
            self.used_gb,
            free_formula,
            used_percent_formula,
            warning_formula,
            self.notes,
            assigned_count_formula,
            deletion_status_formula,
        )


class DatastoreWorksheet:
    """Builds a professional datastore capacity worksheet."""

    sheet_name = "Datastore"
    table_name = "DatastoreCapacityTable"

    def __init__(self, records: list[DatastoreRecord] | None = None) -> None:
        """Initialize the worksheet builder with optional datastore records."""
        self.records = records if records is not None else self._sample_records()
        self.theme = workbook_theme()
        self.header = header_style("datastore_header")
        self.table = table_styles()

    def build(self, workbook: Workbook) -> Worksheet:
        """Create and format the Datastore worksheet."""
        worksheet = self._get_worksheet(workbook)
        self._reset_worksheet(worksheet)
        self._configure_worksheet(worksheet)
        self._write_headers(worksheet)
        self._write_sample_data(worksheet)
        self._apply_table(worksheet)
        self._apply_dropdowns(worksheet)
        self._apply_body_formatting(worksheet)
        self._apply_conditional_formatting(worksheet)
        self._auto_fit_columns(worksheet)
        return worksheet

    def _get_worksheet(self, workbook: Workbook) -> Worksheet:
        """Return the datastore worksheet without creating unrelated sheets."""
        if self.sheet_name in workbook.sheetnames:
            return workbook[self.sheet_name]

        if len(workbook.worksheets) == 1 and workbook.active.max_row == 1:
            worksheet = workbook.active
            worksheet.title = self.sheet_name
            return worksheet

        return workbook.create_sheet(self.sheet_name)

    @staticmethod
    def _reset_worksheet(worksheet: Worksheet) -> None:
        """Clear existing content, tables, and conditional formatting."""
        for merged_range in list(worksheet.merged_cells.ranges):
            worksheet.unmerge_cells(str(merged_range))

        if worksheet.max_row:
            worksheet.delete_rows(1, worksheet.max_row)

        worksheet.tables.clear()
        worksheet.conditional_formatting._cf_rules.clear()

    def _configure_worksheet(self, worksheet: Worksheet) -> None:
        """Apply professional worksheet layout settings."""
        worksheet.sheet_view.showGridLines = False
        worksheet.freeze_panes = "A2"
        worksheet.sheet_properties.tabColor = self.theme.palette.blue
        worksheet.page_setup.orientation = "landscape"
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0
        worksheet.row_dimensions[1].height = 28

    def _write_headers(self, worksheet: Worksheet) -> None:
        """Write datastore capacity headers."""
        for column_index, header in enumerate(DATASTORE_COLUMNS, start=1):
            cell = worksheet.cell(row=1, column=column_index, value=header)
            self._apply_named_style(cell, self.header)

    def _write_sample_data(self, worksheet: Worksheet) -> None:
        """Write sample datastore data and formula columns."""
        for row_index, record in enumerate(self.records, start=2):
            for column_index, value in enumerate(record.as_row(row_index), start=1):
                worksheet.cell(row=row_index, column=column_index, value=value)

    def _apply_table(self, worksheet: Worksheet) -> None:
        """Create an Excel table with filters over the datastore range."""
        table_range = f"A1:K{worksheet.max_row}"
        table = Table(displayName=self.table_name, ref=table_range)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )

        worksheet.add_table(table)

    @staticmethod
    def _apply_dropdowns(worksheet: Worksheet) -> None:
        """Add governed dropdowns to datastore fields."""
        validations = {
            "B2:B1000": ('"VMFS,NFS,vSAN,VVol"', "Select datastore type."),
            "C2:C1000": ("=RefClusters", "Select a governed cluster."),
        }
        for cell_range, (formula, prompt) in validations.items():
            validation = DataValidation(
                type="list",
                formula1=formula,
                allow_blank=True,
            )
            validation.error = "Select a valid value from the dropdown list."
            validation.errorTitle = "Invalid Selection"
            validation.prompt = prompt
            validation.promptTitle = "Governed Dropdown"
            validation.showErrorMessage = True
            validation.showInputMessage = True
            worksheet.add_data_validation(validation)
            validation.add(cell_range)

    def _apply_body_formatting(self, worksheet: Worksheet) -> None:
        """Apply corporate table styles and number formats."""
        numeric_columns = {4, 5, 6, 10}

        for row in worksheet.iter_rows(
            min_row=2,
            max_row=worksheet.max_row,
            min_col=1,
            max_col=len(DATASTORE_COLUMNS),
        ):
            for cell in row:
                style = (
                    self.table["table_numeric"]
                    if cell.column in numeric_columns or cell.column == 7
                    else self.table["table_body"]
                )
                self._apply_named_style(cell, style)

                if cell.column in numeric_columns:
                    cell.number_format = '#,##0'
                elif cell.column == 7:
                    cell.number_format = "0.0%"

        for row_index in range(2, worksheet.max_row + 1):
            worksheet.row_dimensions[row_index].height = 22

    @staticmethod
    def _apply_conditional_formatting(worksheet: Worksheet) -> None:
        """Apply utilization and warning conditional formatting rules."""
        max_row = max(worksheet.max_row, 2)

        worksheet.conditional_formatting.add(
            f"G2:G{max_row}",
            CellIsRule(
                operator="greaterThanOrEqual",
                formula=["0.9"],
                fill=PatternFill("solid", fgColor="F4CCCC"),
                font=Font(color="9C0006", bold=True),
            ),
        )
        worksheet.conditional_formatting.add(
            f"G2:G{max_row}",
            CellIsRule(
                operator="between",
                formula=["0.8", "0.8999"],
                fill=PatternFill("solid", fgColor="FFF2CC"),
                font=Font(color="7F6000", bold=True),
            ),
        )
        worksheet.conditional_formatting.add(
            f"H2:H{max_row}",
            CellIsRule(
                operator="equal",
                formula=['"Critical"'],
                fill=PatternFill("solid", fgColor="F4CCCC"),
                font=Font(color="9C0006", bold=True),
            ),
        )
        worksheet.conditional_formatting.add(
            f"H2:H{max_row}",
            CellIsRule(
                operator="equal",
                formula=['"Warning"'],
                fill=PatternFill("solid", fgColor="FFF2CC"),
                font=Font(color="7F6000", bold=True),
            ),
        )
        worksheet.conditional_formatting.add(
            f"H2:H{max_row}",
            CellIsRule(
                operator="equal",
                formula=['"Healthy"'],
                fill=PatternFill("solid", fgColor="E2F0D9"),
                font=Font(color="006100", bold=True),
            ),
        )
        worksheet.conditional_formatting.add(
            f"K2:K{max_row}",
            CellIsRule(
                operator="equal",
                formula=['"Assigned to VM - must not delete"'],
                fill=PatternFill("solid", fgColor="F4CCCC"),
                font=Font(color="9C0006", bold=True),
            ),
        )
        worksheet.conditional_formatting.add(
            f"K2:K{max_row}",
            CellIsRule(
                operator="equal",
                formula=['"Safe to delete"'],
                fill=PatternFill("solid", fgColor="E2F0D9"),
                font=Font(color="006100", bold=True),
            ),
        )

    @staticmethod
    def _apply_named_style(cell: Cell, named_style: Any) -> None:
        """Copy a named style's visual attributes onto a cell."""
        cell.font = copy(named_style.font)
        cell.fill = copy(named_style.fill)
        cell.border = copy(named_style.border)
        cell.alignment = copy(named_style.alignment)
        cell.number_format = named_style.number_format

    @staticmethod
    def _auto_fit_columns(worksheet: Worksheet) -> None:
        """Set practical column widths based on visible cell values."""
        for column_cells in worksheet.columns:
            max_length = 0
            column_letter = column_cells[0].column_letter

            for cell in column_cells:
                if cell.value is None:
                    continue
                max_length = max(max_length, len(str(cell.value)))

            worksheet.column_dimensions[column_letter].width = min(
                max(max_length + 2, 12),
                36,
            )

    @staticmethod
    def _sample_records() -> list[DatastoreRecord]:
        """Return representative datastore capacity rows."""
        return [
            DatastoreRecord(
                datastore="DS-PRD-01",
                datastore_type="VMFS",
                cluster="Cluster-Production",
                capacity_gb=12000,
                used_gb=9300,
                notes="Primary production datastore",
            ),
            DatastoreRecord(
                datastore="DS-PRD-02",
                datastore_type="vSAN",
                cluster="Cluster-Production",
                capacity_gb=18000,
                used_gb=16550,
                notes="Critical utilization threshold exceeded",
            ),
            DatastoreRecord(
                datastore="DS-DEV-01",
                datastore_type="NFS",
                cluster="Cluster-Development",
                capacity_gb=8000,
                used_gb=4200,
                notes="Healthy development capacity",
            ),
            DatastoreRecord(
                datastore="DS-TST-01",
                datastore_type="VMFS",
                cluster="Cluster-Test",
                capacity_gb=6000,
                used_gb=4850,
                notes="Monitor growth trend",
            ),
        ]


