"""Virtualization license worksheet generation."""

from __future__ import annotations

from copy import copy
from dataclasses import dataclass
from datetime import date
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from .styles import header_style, table_styles, workbook_theme


LICENSE_COLUMNS: tuple[str, ...] = (
    "Product",
    "Edition",
    "License Key",
    "Total CPUs",
    "Used CPUs",
    "Available CPUs",
    "Expiration Date",
    "Support Level",
    "Status",
    "Notes",
)


@dataclass(frozen=True)
class LicenseRecord:
    """Virtualization license register row."""

    product: str
    edition: str
    license_key: str
    total_cpus: int
    used_cpus: int
    expiration_date: date | str
    support_level: str
    notes: str

    def as_row(self, row_index: int) -> tuple[Any, ...]:
        """Return the record in worksheet order with formulas."""
        available_formula = f"=D{row_index}-E{row_index}"
        status_formula = (
            f'=IF(F{row_index}<0,"Overallocated",'
            f'IF(G{row_index}="","Perpetual",'
            f'IF(G{row_index}<TODAY(),"Expired","Active")))'
        )
        return (
            self.product,
            self.edition,
            self.license_key,
            self.total_cpus,
            self.used_cpus,
            available_formula,
            self.expiration_date,
            self.support_level,
            status_formula,
            self.notes,
        )


class LicenseWorksheet:
    """Builds a Virtualization license register worksheet."""

    sheet_name = "License"
    table_name = "LicenseRegisterTable"

    def __init__(self, records: list[LicenseRecord] | None = None) -> None:
        """Initialize the worksheet builder."""
        self.records = records if records is not None else self._sample_records()
        self.theme = workbook_theme()
        self.header = header_style("license_header")
        self.table = table_styles()

    def build(self, workbook: Workbook) -> Worksheet:
        """Create and format the License worksheet."""
        worksheet = self._get_worksheet(workbook)
        self._reset_worksheet(worksheet)
        self._configure_worksheet(worksheet)
        self._write_headers(worksheet)
        self._write_sample_data(worksheet)
        self._apply_table(worksheet)
        self._apply_body_formatting(worksheet)
        self._apply_conditional_formatting(worksheet)
        self._auto_fit_columns(worksheet)
        return worksheet

    def _get_worksheet(self, workbook: Workbook) -> Worksheet:
        """Return the License worksheet."""
        if self.sheet_name in workbook.sheetnames:
            return workbook[self.sheet_name]
        return workbook.create_sheet(self.sheet_name)

    @staticmethod
    def _reset_worksheet(worksheet: Worksheet) -> None:
        """Clear worksheet content and rules."""
        for merged_range in list(worksheet.merged_cells.ranges):
            worksheet.unmerge_cells(str(merged_range))
        if worksheet.max_row:
            worksheet.delete_rows(1, worksheet.max_row)
        worksheet.tables.clear()
        worksheet.conditional_formatting._cf_rules.clear()

    def _configure_worksheet(self, worksheet: Worksheet) -> None:
        """Apply worksheet layout settings."""
        worksheet.sheet_view.showGridLines = False
        worksheet.freeze_panes = "A2"
        worksheet.sheet_properties.tabColor = self.theme.palette.navy
        worksheet.page_setup.orientation = "landscape"
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0
        worksheet.row_dimensions[1].height = 28

    def _write_headers(self, worksheet: Worksheet) -> None:
        """Write license register headers."""
        for column_index, header in enumerate(LICENSE_COLUMNS, start=1):
            cell = worksheet.cell(row=1, column=column_index, value=header)
            self._apply_named_style(cell, self.header)

    def _write_sample_data(self, worksheet: Worksheet) -> None:
        """Write sample license data."""
        for row_index, record in enumerate(self.records, start=2):
            for column_index, value in enumerate(record.as_row(row_index), start=1):
                worksheet.cell(row=row_index, column=column_index, value=value)

    def _apply_table(self, worksheet: Worksheet) -> None:
        """Create an Excel table with filters."""
        table_range = f"A1:J{worksheet.max_row}"
        table = Table(displayName=self.table_name, ref=table_range)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        worksheet.add_table(table)

    def _apply_body_formatting(self, worksheet: Worksheet) -> None:
        """Apply table styling and number formats."""
        numeric_columns = {4, 5, 6}
        for row in worksheet.iter_rows(
            min_row=2,
            max_row=worksheet.max_row,
            min_col=1,
            max_col=len(LICENSE_COLUMNS),
        ):
            for cell in row:
                style = (
                    self.table["table_numeric"]
                    if cell.column in numeric_columns
                    else self.table["table_body"]
                )
                self._apply_named_style(cell, style)
                if cell.column == 7:
                    cell.number_format = "yyyy-mm-dd"

    @staticmethod
    def _apply_conditional_formatting(worksheet: Worksheet) -> None:
        """Highlight license status values."""
        max_row = max(worksheet.max_row, 2)
        rules = (
            ('"Active"', "E2F0D9", "006100"),
            ('"Perpetual"', "D9EAF7", "1F4E78"),
            ('"Expired"', "F4CCCC", "9C0006"),
            ('"Overallocated"', "F4CCCC", "9C0006"),
        )
        for formula, fill_color, font_color in rules:
            worksheet.conditional_formatting.add(
                f"I2:I{max_row}",
                CellIsRule(
                    operator="equal",
                    formula=[formula],
                    fill=PatternFill("solid", fgColor=fill_color),
                    font=Font(color=font_color, bold=True),
                ),
            )

    @staticmethod
    def _apply_named_style(cell: Cell, named_style: Any) -> None:
        """Copy a named style's visual attributes onto a cell."""
        cell.font = copy(named_style.font)
        cell.fill = copy(named_style.fill)
        cell.border = copy(named_style.border)
        cell.alignment = copy(named_style.alignment)
        cell.number_format = named_style.number_format

    @staticmethod
    def _auto_fit_columns(worksheet: Worksheet) -> None:
        """Set practical column widths."""
        for column_cells in worksheet.columns:
            max_length = 0
            column_letter = column_cells[0].column_letter
            for cell in column_cells:
                if cell.value is not None:
                    max_length = max(max_length, len(str(cell.value)))
            worksheet.column_dimensions[column_letter].width = min(
                max(max_length + 2, 12),
                34,
            )

    @staticmethod
    def _sample_records() -> list[LicenseRecord]:
        """Return representative Virtualization license rows."""
        return [
            LicenseRecord(
                product="Virtualization Hypervisor Suite",
                edition="Enterprise Plus",
                license_key="XXXXX-XXXXX-XXXXX-XXXXX-XXXXX",
                total_cpus=24,
                used_cpus=20,
                expiration_date="",
                support_level="Production Support",
                notes="Primary production entitlement",
            ),
            LicenseRecord(
                product="Virtualization Management Server",
                edition="Standard",
                license_key="YYYYY-YYYYY-YYYYY-YYYYY-YYYYY",
                total_cpus=1,
                used_cpus=1,
                expiration_date=date(2027, 6, 30),
                support_level="Production Support",
                notes="Central management license",
            ),
        ]


