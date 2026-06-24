"""Host Inventory worksheet generation.

This module creates only the ``Host Inventory`` worksheet.
"""

from __future__ import annotations

from copy import copy
from dataclasses import dataclass
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from .styles import header_style, table_styles, workbook_theme


HOST_INVENTORY_COLUMNS: tuple[str, ...] = (
    "Hostname",
    "Version",
    "Build",
    "CPU",
    "Memory",
    "NICs",
    "VM Count",
    "Datastores",
    "Cluster",
    "Management IP",
    "License",
    "Health",
    "Notes",
    "Assigned VM Count",
    "Deletion Status",
)


@dataclass(frozen=True)
class HostRecord:
    """Sample host inventory row."""

    hostname: str
    version: str
    build: str
    cpu: str
    memory: str
    nics: int
    vm_count: int
    datastores: str
    cluster: str
    management_ip: str
    license: str
    health: str
    notes: str

    def as_row(self, row_index: int) -> tuple[Any, ...]:
        """Return the record in worksheet column order."""
        assigned_count_formula = (
            f"=COUNTIF('VM Inventory'!$H:$H,A{row_index})"
        )
        deletion_status_formula = (
            f'=IF(N{row_index}>0,'
            '"Assigned to VM - must not delete","Safe to delete")'
        )
        return (
            self.hostname,
            self.version,
            self.build,
            self.cpu,
            self.memory,
            self.nics,
            self.vm_count,
            self.datastores,
            self.cluster,
            self.management_ip,
            self.license,
            self.health,
            self.notes,
            assigned_count_formula,
            deletion_status_formula,
        )


class HostInventoryWorksheet:
    """Builds a professional host inventory worksheet."""

    sheet_name = "Host Inventory"
    table_name = "HostInventoryTable"

    def __init__(self, records: list[HostRecord] | None = None) -> None:
        """Initialize the worksheet builder with optional host records."""
        self.records = records if records is not None else self._sample_records()
        self.theme = workbook_theme()
        self.header = header_style("host_inventory_header")
        self.table = table_styles()

    def build(self, workbook: Workbook) -> Worksheet:
        """Create and format the Host Inventory worksheet."""
        worksheet = self._get_worksheet(workbook)
        self._reset_worksheet(worksheet)
        self._configure_worksheet(worksheet)
        self._write_headers(worksheet)
        self._write_sample_data(worksheet)
        self._apply_table(worksheet)
        self._apply_dropdowns(worksheet)
        self._apply_body_formatting(worksheet)
        self._apply_conditional_formatting(worksheet)
        self._auto_fit_columns(worksheet)
        return worksheet

    def _get_worksheet(self, workbook: Workbook) -> Worksheet:
        """Return the worksheet without creating unrelated sheets."""
        if self.sheet_name in workbook.sheetnames:
            return workbook[self.sheet_name]

        if len(workbook.worksheets) == 1 and workbook.active.max_row == 1:
            worksheet = workbook.active
            worksheet.title = self.sheet_name
            return worksheet

        return workbook.create_sheet(self.sheet_name)

    @staticmethod
    def _reset_worksheet(worksheet: Worksheet) -> None:
        """Clear existing content and table definitions before rebuilding."""
        for merged_range in list(worksheet.merged_cells.ranges):
            worksheet.unmerge_cells(str(merged_range))

        if worksheet.max_row:
            worksheet.delete_rows(1, worksheet.max_row)

        worksheet.tables.clear()
        worksheet.data_validations.dataValidation = []
        worksheet.conditional_formatting._cf_rules.clear()

    def _configure_worksheet(self, worksheet: Worksheet) -> None:
        """Apply sheet-level professional layout settings."""
        worksheet.sheet_view.showGridLines = False
        worksheet.freeze_panes = "A2"
        worksheet.sheet_properties.tabColor = self.theme.palette.navy
        worksheet.page_setup.orientation = "landscape"
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0
        worksheet.row_dimensions[1].height = 28

    def _write_headers(self, worksheet: Worksheet) -> None:
        """Write host inventory headers."""
        for column_index, header in enumerate(HOST_INVENTORY_COLUMNS, start=1):
            cell = worksheet.cell(row=1, column=column_index, value=header)
            self._apply_named_style(cell, self.header)

    def _write_sample_data(self, worksheet: Worksheet) -> None:
        """Write representative host sample data."""
        for row_index, record in enumerate(self.records, start=2):
            for column_index, value in enumerate(record.as_row(row_index), start=1):
                worksheet.cell(row=row_index, column=column_index, value=value)

    def _apply_table(self, worksheet: Worksheet) -> None:
        """Create an Excel table with filters over the host inventory range."""
        table_range = f"A1:O{worksheet.max_row}"
        table = Table(displayName=self.table_name, ref=table_range)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )

        worksheet.add_table(table)

    @staticmethod
    def _apply_dropdowns(worksheet: Worksheet) -> None:
        """Add governed dropdowns to host inventory fields."""
        validations = {
            "K2:K1000": (
                '"Enterprise Plus,Standard,Management Server Standard"',
                "Select a license edition.",
            ),
            "L2:L1000": (
                '"Healthy,Warning,Degraded,Critical,Disconnected"',
                "Select health.",
            ),
        }
        for cell_range, (formula, prompt) in validations.items():
            validation = DataValidation(
                type="list",
                formula1=formula,
                allow_blank=True,
            )
            validation.error = "Select a valid value from the dropdown list."
            validation.errorTitle = "Invalid Selection"
            validation.prompt = prompt
            validation.promptTitle = "Governed Dropdown"
            validation.showErrorMessage = True
            validation.showInputMessage = True
            worksheet.add_data_validation(validation)
            validation.add(cell_range)

    def _apply_body_formatting(self, worksheet: Worksheet) -> None:
        """Apply corporate table styling and health status colors."""
        numeric_columns = {6, 7, 14}

        for row in worksheet.iter_rows(
            min_row=2,
            max_row=worksheet.max_row,
            min_col=1,
            max_col=len(HOST_INVENTORY_COLUMNS),
        ):
            for cell in row:
                style = (
                    self.table["table_numeric"]
                    if cell.column in numeric_columns
                    else self.table["table_body"]
                )
                self._apply_named_style(cell, style)

                if cell.column == 12:
                    self._apply_health_fill(cell)

        for row_index in range(2, worksheet.max_row + 1):
            worksheet.row_dimensions[row_index].height = 22

    @staticmethod
    def _apply_conditional_formatting(worksheet: Worksheet) -> None:
        """Flag hosts that are still assigned to one or more VMs."""
        max_row = max(worksheet.max_row, 2)
        worksheet.conditional_formatting.add(
            f"O2:O{max_row}",
            CellIsRule(
                operator="equal",
                formula=['"Assigned to VM - must not delete"'],
                fill=PatternFill("solid", fgColor="F4CCCC"),
                font=Font(color="9C0006", bold=True),
            ),
        )
        worksheet.conditional_formatting.add(
            f"O2:O{max_row}",
            CellIsRule(
                operator="equal",
                formula=['"Safe to delete"'],
                fill=PatternFill("solid", fgColor="E2F0D9"),
                font=Font(color="006100", bold=True),
            ),
        )

    @staticmethod
    def _apply_health_fill(cell: Cell) -> None:
        """Apply subtle status color to the Health column."""
        value = str(cell.value).lower()

        if value == "healthy":
            cell.fill = PatternFill("solid", fgColor="E2F0D9")
            cell.font = Font(name="Calibri", size=11, bold=True, color="006100")
        elif value in {"warning", "degraded"}:
            cell.fill = PatternFill("solid", fgColor="FFF2CC")
            cell.font = Font(name="Calibri", size=11, bold=True, color="7F6000")
        elif value in {"critical", "disconnected"}:
            cell.fill = PatternFill("solid", fgColor="F4CCCC")
            cell.font = Font(name="Calibri", size=11, bold=True, color="9C0006")

    @staticmethod
    def _apply_named_style(cell: Cell, named_style: Any) -> None:
        """Copy a named style's visual attributes onto a cell."""
        cell.font = copy(named_style.font)
        cell.fill = copy(named_style.fill)
        cell.border = copy(named_style.border)
        cell.alignment = copy(named_style.alignment)
        cell.number_format = named_style.number_format

    @staticmethod
    def _auto_fit_columns(worksheet: Worksheet) -> None:
        """Set practical column widths based on visible cell values."""
        for column_cells in worksheet.columns:
            max_length = 0
            column_letter = column_cells[0].column_letter

            for cell in column_cells:
                if cell.value is None:
                    continue
                max_length = max(max_length, len(str(cell.value)))

            worksheet.column_dimensions[column_letter].width = min(
                max(max_length + 2, 12),
                36,
            )

    @staticmethod
    def _sample_records() -> list[HostRecord]:
        """Return representative host inventory rows."""
        return [
            HostRecord(
                hostname="host-prd-01.company.local",
                version="Generic Hypervisor 2026",
                build="22380479",
                cpu="2 x Intel Xeon Gold 6348, 56 cores",
                memory="1024 GB",
                nics=8,
                vm_count=42,
                datastores="DS-PRD-01, DS-PRD-02",
                cluster="Cluster-Production",
                management_ip="10.10.1.21",
                license="Enterprise Plus",
                health="Healthy",
                notes="Primary production compute host",
            ),
            HostRecord(
                hostname="host-prd-02.company.local",
                version="Generic Hypervisor 2026",
                build="22380479",
                cpu="2 x Intel Xeon Gold 6348, 56 cores",
                memory="1024 GB",
                nics=8,
                vm_count=39,
                datastores="DS-PRD-01, DS-PRD-02",
                cluster="Cluster-Production",
                management_ip="10.10.1.22",
                license="Enterprise Plus",
                health="Warning",
                notes="NIC redundancy warning under review",
            ),
            HostRecord(
                hostname="host-dev-01.company.local",
                version="Generic Hypervisor 2025",
                build="21495797",
                cpu="2 x Intel Xeon Silver 4314, 32 cores",
                memory="512 GB",
                nics=6,
                vm_count=24,
                datastores="DS-DEV-01, DS-DEV-02",
                cluster="Cluster-Development",
                management_ip="10.20.1.31",
                license="Standard",
                health="Healthy",
                notes="Development workload host",
            ),
        ]


