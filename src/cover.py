"""Premium cover page generation for Virtualization Administration Toolkit.

This module is responsible only for creating and formatting the cover
worksheet. It does not create or manage any other workbook worksheets.
"""

from __future__ import annotations

from dataclasses import dataclass

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, GradientFill, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

from .styles import CorporatePalette, corporate_color_palette


@dataclass(frozen=True)
class CoverMetadata:
    """Metadata displayed on the cover worksheet."""

    title: str = "Enterprise Virtualization\nAdministration Toolkit"
    subtitle: str = "Professional Infrastructure Reporting Workbook"
    version: str = "Version 1.0.0"
    author: str = "Author: <Author Name>"
    company: str = "Company: <Company Name>"
    footer: str = "(c) 2026 <Company Name>. All rights reserved."
    brand_band: str = "ENTERPRISE VIRTUALIZATION ADMINISTRATION TOOLKIT"


class CoverWorksheet:
    """Builds the premium cover worksheet for the report workbook."""

    sheet_name = "Cover"

    def __init__(
        self,
        metadata: CoverMetadata | None = None,
        palette: CorporatePalette | None = None,
    ) -> None:
        """Initialize cover worksheet settings."""
        self.metadata = metadata or CoverMetadata()
        self.palette = palette or corporate_color_palette()

    def build(self, workbook: Workbook) -> Worksheet:
        """Create and format the cover worksheet only.

        A blank default worksheet is renamed and reused when possible so that a
        new workbook contains only the cover worksheet after this method runs.
        """
        worksheet = self._get_cover_worksheet(workbook)
        self._reset_cover_sheet(worksheet)
        self._configure_page(worksheet)
        self._apply_background(worksheet)
        self._add_brand_band(worksheet)
        self._add_title_block(worksheet)
        self._add_metadata_panel(worksheet)
        self._add_feature_icons(worksheet)
        self._add_footer(worksheet)
        return worksheet

    def _get_cover_worksheet(self, workbook: Workbook) -> Worksheet:
        """Return the cover worksheet without creating unrelated sheets."""
        if self.sheet_name in workbook.sheetnames:
            return workbook[self.sheet_name]

        if len(workbook.worksheets) == 1 and workbook.active.max_row == 1:
            worksheet = workbook.active
            worksheet.title = self.sheet_name
            return worksheet

        return workbook.create_sheet(self.sheet_name, 0)

    @staticmethod
    def _reset_cover_sheet(worksheet: Worksheet) -> None:
        """Clear existing cover content before rebuilding the page."""
        for merged_range in list(worksheet.merged_cells.ranges):
            worksheet.unmerge_cells(str(merged_range))

        for row in worksheet.iter_rows():
            for cell in row:
                cell.value = None
                cell.style = "Normal"

    def _configure_page(self, worksheet: Worksheet) -> None:
        """Set page sizing, margins, and print-friendly configuration."""
        worksheet.sheet_view.showGridLines = False
        worksheet.freeze_panes = None
        worksheet.page_setup.orientation = "landscape"
        worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
        worksheet.page_margins.left = 0.25
        worksheet.page_margins.right = 0.25
        worksheet.page_margins.top = 0.25
        worksheet.page_margins.bottom = 0.25

        column_widths = {
            "A": 4,
            "B": 14,
            "C": 14,
            "D": 18,
            "E": 18,
            "F": 18,
            "G": 18,
            "H": 18,
            "I": 14,
            "J": 14,
            "K": 4,
        }
        for column, width in column_widths.items():
            worksheet.column_dimensions[column].width = width

        for row_index in range(1, 38):
            worksheet.row_dimensions[row_index].height = 24

    def _apply_background(self, worksheet: Worksheet) -> None:
        """Apply the premium blue cover background."""
        background_fill = GradientFill(
            type="linear",
            degree=0,
            stop=(self.palette.navy, "244D73", self.palette.blue),
        )

        for row in worksheet["A1:K37"]:
            for cell in row:
                cell.fill = background_fill

        accent_fill = PatternFill("solid", fgColor=self.palette.light_blue)
        for row in worksheet["A1:A37"]:
            row[0].fill = accent_fill

    def _add_brand_band(self, worksheet: Worksheet) -> None:
        """Add the top corporate branding band."""
        worksheet.merge_cells("B2:J3")
        brand_cell = worksheet["B2"]
        brand_cell.value = self.metadata.brand_band
        brand_cell.font = Font(
            name="Calibri",
            size=15,
            bold=True,
            color=self.palette.white,
        )
        brand_cell.alignment = Alignment(horizontal="left", vertical="center")

        worksheet.merge_cells("B4:J4")
        divider = worksheet["B4"]
        divider.value = ""
        divider.fill = PatternFill("solid", fgColor=self.palette.light_blue)

    def _add_title_block(self, worksheet: Worksheet) -> None:
        """Add the large premium title and subtitle."""
        worksheet.merge_cells("B8:J11")
        title_cell = worksheet["B8"]
        title_cell.value = self.metadata.title
        title_cell.font = Font(
            name="Calibri Light",
            size=34,
            bold=True,
            color=self.palette.white,
        )
        title_cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

        worksheet.merge_cells("C13:I14")
        subtitle_cell = worksheet["C13"]
        subtitle_cell.value = self.metadata.subtitle
        subtitle_cell.font = Font(
            name="Calibri",
            size=15,
            italic=True,
            color=self.palette.light_blue,
        )
        subtitle_cell.alignment = Alignment(horizontal="center", vertical="center")

    def _add_metadata_panel(self, worksheet: Worksheet) -> None:
        """Add version, author, and company placeholders."""
        panel_fill = PatternFill("solid", fgColor="EAF3F8")
        panel_border = Border(
            left=Side(style="thin", color=self.palette.white),
            right=Side(style="thin", color=self.palette.white),
            top=Side(style="thin", color=self.palette.white),
            bottom=Side(style="thin", color=self.palette.white),
        )

        for row in worksheet["D17:H21"]:
            for cell in row:
                cell.fill = panel_fill
                cell.border = panel_border

        worksheet.merge_cells("D17:H17")
        panel_title = worksheet["D17"]
        panel_title.value = "Report Information"
        panel_title.font = Font(
            name="Calibri",
            size=12,
            bold=True,
            color=self.palette.white,
        )
        panel_title.fill = PatternFill("solid", fgColor=self.palette.navy)
        panel_title.alignment = Alignment(horizontal="center", vertical="center")

        metadata_rows = {
            "E18": f"{self.metadata.version}",
            "E19": f"{self.metadata.author}",
            "E20": f"{self.metadata.company}",
        }

        for coordinate, value in metadata_rows.items():
            cell = worksheet[coordinate]
            cell.value = value
            cell.font = Font(
                name="Calibri",
                size=13,
                bold=True,
                color=self.palette.navy,
            )
            cell.alignment = Alignment(horizontal="left", vertical="center")

    def _add_feature_icons(self, worksheet: Worksheet) -> None:
        """Add clickable navigation tiles for workbook sheets."""
        navigation_items = (
            ("B23:D24", "Dashboard", "Executive Dashboard"),
            ("E23:G24", "VM Inventory", "VM Inventory"),
            ("H23:J24", "Hosts", "Host Inventory"),
            ("B26:D27", "Datastores", "Datastore"),
            ("E26:G27", "Capacity", "Capacity Planner"),
            ("H26:J27", "Snapshots", "Snapshot Tracker"),
            ("B29:D30", "Backups", "Backup Register"),
            ("E29:G30", "Licensing", "License"),
            ("H29:J30", "Reference Data", "Reference Data"),
            ("E32:G33", "Instructions", "Instructions"),
        )

        for cell_range, label, sheet_name in navigation_items:
            self._add_navigation_tile(worksheet, cell_range, label, sheet_name)

    def _add_navigation_tile(
        self,
        worksheet: Worksheet,
        cell_range: str,
        label: str,
        sheet_name: str,
    ) -> None:
        """Create one clickable cover navigation tile."""
        worksheet.merge_cells(cell_range)
        top_left = cell_range.split(":")[0]
        cell = worksheet[top_left]
        cell.value = label
        cell.hyperlink = f"#'{sheet_name}'!A1"
        cell.style = "Hyperlink"
        cell.font = Font(
            name="Calibri",
            size=12,
            bold=True,
            color=self.palette.white,
            underline="single",
        )
        cell.fill = PatternFill("solid", fgColor="244D73")
        cell.alignment = Alignment(horizontal="center", vertical="center")

        tile_border = Border(
            left=Side(style="thin", color=self.palette.light_blue),
            right=Side(style="thin", color=self.palette.light_blue),
            top=Side(style="thin", color=self.palette.light_blue),
            bottom=Side(style="thin", color=self.palette.light_blue),
        )
        for row in worksheet[cell_range]:
            for tile_cell in row:
                tile_cell.fill = PatternFill("solid", fgColor="244D73")
                tile_cell.border = tile_border

    def _add_footer(self, worksheet: Worksheet) -> None:
        """Add the professional footer text."""
        worksheet.merge_cells("B35:J36")
        footer_cell = worksheet["B35"]
        footer_cell.value = self.metadata.footer
        footer_cell.font = Font(
            name="Calibri",
            size=10,
            italic=True,
            color=self.palette.light_blue,
        )
        footer_cell.alignment = Alignment(horizontal="center", vertical="center")


