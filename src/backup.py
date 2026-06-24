"""Backup Register worksheet generation.

This module creates only the ``Backup Register`` worksheet. It includes a
backup register table and an executive dashboard section on the same sheet.
"""

from __future__ import annotations

from copy import copy
from dataclasses import dataclass
from datetime import date
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from .styles import header_style, table_styles, workbook_theme


BACKUP_COLUMNS: tuple[str, ...] = (
    "Backup Job",
    "Schedule",
    "Repository",
    "Retention",
    "Last Backup",
    "Status",
    "Successful Runs",
    "Total Runs",
    "Success %",
    "Restore Test",
    "Restore Test Date",
    "Protected VMs",
    "Owner",
    "Notes",
)

TABLE_START_ROW = 8


@dataclass(frozen=True)
class BackupRecord:
    """Sample backup register row."""

    backup_job: str
    schedule: str
    repository: str
    retention: str
    last_backup: date
    status: str
    successful_runs: int
    total_runs: int
    restore_test: str
    restore_test_date: date | str
    protected_vms: int
    owner: str
    notes: str

    def as_row(self, row_index: int) -> tuple[Any, ...]:
        """Return the record in worksheet order with success formula."""
        success_formula = f"=IFERROR(G{row_index}/H{row_index},0)"
        return (
            self.backup_job,
            self.schedule,
            self.repository,
            self.retention,
            self.last_backup,
            self.status,
            self.successful_runs,
            self.total_runs,
            success_formula,
            self.restore_test,
            self.restore_test_date,
            self.protected_vms,
            self.owner,
            self.notes,
        )


class BackupWorksheet:
    """Builds a professional Backup Register worksheet."""

    sheet_name = "Backup Register"
    table_name = "BackupRegisterTable"

    def __init__(self, records: list[BackupRecord] | None = None) -> None:
        """Initialize the worksheet builder with optional backup records."""
        self.records = records if records is not None else self._sample_records()
        self.theme = workbook_theme()
        self.palette = self.theme.palette
        self.header = header_style("backup_register_header")
        self.table = table_styles()

    def build(self, workbook: Workbook) -> Worksheet:
        """Create and format the Backup Register worksheet."""
        worksheet = self._get_worksheet(workbook)
        self._reset_worksheet(worksheet)
        self._configure_worksheet(worksheet)
        self._add_dashboard(worksheet)
        self._write_headers(worksheet)
        self._write_sample_data(worksheet)
        self._apply_table(worksheet)
        self._apply_dropdowns(worksheet)
        self._apply_body_formatting(worksheet)
        self._apply_conditional_formatting(worksheet)
        self._auto_fit_columns(worksheet)
        return worksheet

    def _get_worksheet(self, workbook: Workbook) -> Worksheet:
        """Return the Backup Register worksheet without adding other sheets."""
        if self.sheet_name in workbook.sheetnames:
            return workbook[self.sheet_name]

        if len(workbook.worksheets) == 1 and workbook.active.max_row == 1:
            worksheet = workbook.active
            worksheet.title = self.sheet_name
            return worksheet

        return workbook.create_sheet(self.sheet_name)

    @staticmethod
    def _reset_worksheet(worksheet: Worksheet) -> None:
        """Clear content, merged cells, tables, and formatting rules."""
        for merged_range in list(worksheet.merged_cells.ranges):
            worksheet.unmerge_cells(str(merged_range))

        if worksheet.max_row:
            worksheet.delete_rows(1, worksheet.max_row)

        worksheet.tables.clear()
        worksheet.conditional_formatting._cf_rules.clear()

    def _configure_worksheet(self, worksheet: Worksheet) -> None:
        """Apply professional worksheet layout settings."""
        worksheet.sheet_view.showGridLines = False
        worksheet.freeze_panes = f"A{TABLE_START_ROW + 1}"
        worksheet.sheet_properties.tabColor = self.palette.green
        worksheet.page_setup.orientation = "landscape"
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0

        for row_index in range(1, TABLE_START_ROW):
            worksheet.row_dimensions[row_index].height = 25
        worksheet.row_dimensions[TABLE_START_ROW].height = 28

    def _add_dashboard(self, worksheet: Worksheet) -> None:
        """Add dashboard title and KPI cards above the backup register."""
        worksheet.merge_cells("A1:N2")
        title_cell = worksheet["A1"]
        title_cell.value = "Backup Register Dashboard"
        title_cell.font = Font(
            name="Calibri Light",
            size=26,
            bold=True,
            color=self.palette.white,
        )
        title_cell.fill = PatternFill("solid", fgColor=self.palette.navy)
        title_cell.alignment = Alignment(horizontal="center", vertical="center")

        cards = (
            ("A4:C6", "Total Jobs", "=COUNTA(A9:A1000)", self.palette.navy),
            ("D4:F6", "Successful", '=COUNTIF(F9:F1000,"Success")', self.palette.green),
            ("G4:I6", "Failed", '=COUNTIF(F9:F1000,"Failed")', self.palette.red),
            (
                "J4:L6",
                "Restore Tests",
                '=COUNTIF(J9:J1000,"Passed")',
                self.palette.blue,
            ),
        )
        for cell_range, title, formula, color in cards:
            self._add_kpi_card(worksheet, cell_range, title, formula, color)

    def _add_kpi_card(
        self,
        worksheet: Worksheet,
        cell_range: str,
        title: str,
        formula: str,
        color: str,
    ) -> None:
        """Render a merged KPI card for the backup dashboard."""
        min_col, min_row, max_col, max_row = range_boundaries(cell_range)
        title_range = (
            f"{get_column_letter(min_col)}{min_row}:"
            f"{get_column_letter(max_col)}{min_row}"
        )
        value_range = (
            f"{get_column_letter(min_col)}{min_row + 1}:"
            f"{get_column_letter(max_col)}{max_row}"
        )
        worksheet.merge_cells(title_range)
        worksheet.merge_cells(value_range)

        title_cell = worksheet.cell(row=min_row, column=min_col)
        title_cell.value = title
        title_cell.font = Font(name="Calibri", size=13, bold=True, color=color)
        title_cell.fill = PatternFill("solid", fgColor=self.palette.white)
        title_cell.alignment = Alignment(horizontal="center", vertical="center")

        value_cell = worksheet.cell(row=min_row + 1, column=min_col)
        value_cell.value = formula
        value_cell.font = Font(name="Calibri", size=20, bold=True, color=color)
        value_cell.fill = PatternFill("solid", fgColor=self.palette.white)
        value_cell.alignment = Alignment(horizontal="center", vertical="center")

        border = Border(
            left=Side(style="medium", color=color),
            right=Side(style="medium", color=color),
            top=Side(style="medium", color=color),
            bottom=Side(style="medium", color=color),
        )
        for row in worksheet[cell_range]:
            for card_cell in row:
                card_cell.border = border

    def _write_headers(self, worksheet: Worksheet) -> None:
        """Write Backup Register table headers."""
        for column_index, header in enumerate(BACKUP_COLUMNS, start=1):
            cell = worksheet.cell(
                row=TABLE_START_ROW,
                column=column_index,
                value=header,
            )
            self._apply_named_style(cell, self.header)

    def _write_sample_data(self, worksheet: Worksheet) -> None:
        """Write representative backup job records."""
        for row_offset, record in enumerate(self.records, start=1):
            row_index = TABLE_START_ROW + row_offset
            for column_index, value in enumerate(record.as_row(row_index), start=1):
                worksheet.cell(row=row_index, column=column_index, value=value)

    def _apply_table(self, worksheet: Worksheet) -> None:
        """Create an Excel table with filters over backup register data."""
        table_range = f"A{TABLE_START_ROW}:N{worksheet.max_row}"
        table = Table(displayName=self.table_name, ref=table_range)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )

        worksheet.add_table(table)

    @staticmethod
    def _apply_dropdowns(worksheet: Worksheet) -> None:
        """Add governed dropdowns to backup register fields."""
        validations = {
            "A9:A1000": ("=RefBackupJobs", "Select a governed backup job."),
            "C9:C1000": ("=RefRepositories", "Select a governed repository."),
            "F9:F1000": ('"Success,Warning,Failed"', "Select backup status."),
            "J9:J1000": ('"Passed,Pending,Failed"', "Select restore test status."),
            "M9:M1000": ("=RefOwners", "Select a governed owner."),
        }
        for cell_range, (formula, prompt) in validations.items():
            validation = DataValidation(
                type="list",
                formula1=formula,
                allow_blank=True,
            )
            validation.error = "Select a valid value from the dropdown list."
            validation.errorTitle = "Invalid Selection"
            validation.prompt = prompt
            validation.promptTitle = "Governed Dropdown"
            validation.showErrorMessage = True
            validation.showInputMessage = True
            worksheet.add_data_validation(validation)
            validation.add(cell_range)

    def _apply_body_formatting(self, worksheet: Worksheet) -> None:
        """Apply corporate table styling and number formats."""
        numeric_columns = {7, 8, 12}
        date_columns = {5, 11}

        for row in worksheet.iter_rows(
            min_row=TABLE_START_ROW + 1,
            max_row=worksheet.max_row,
            min_col=1,
            max_col=len(BACKUP_COLUMNS),
        ):
            for cell in row:
                style = (
                    self.table["table_numeric"]
                    if cell.column in numeric_columns or cell.column == 9
                    else self.table["table_body"]
                )
                self._apply_named_style(cell, style)

                if cell.column in date_columns:
                    cell.number_format = "yyyy-mm-dd"
                elif cell.column == 9:
                    cell.number_format = "0.0%"

        for row_index in range(TABLE_START_ROW + 1, worksheet.max_row + 1):
            worksheet.row_dimensions[row_index].height = 22

    @staticmethod
    def _apply_conditional_formatting(worksheet: Worksheet) -> None:
        """Apply status, success rate, and restore-test warning colors."""
        max_row = max(worksheet.max_row, TABLE_START_ROW + 1)

        status_rules = (
            ('"Success"', "E2F0D9", "006100"),
            ('"Warning"', "FFF2CC", "7F6000"),
            ('"Failed"', "F4CCCC", "9C0006"),
        )
        for formula, fill_color, font_color in status_rules:
            worksheet.conditional_formatting.add(
                f"F9:F{max_row}",
                CellIsRule(
                    operator="equal",
                    formula=[formula],
                    fill=PatternFill("solid", fgColor=fill_color),
                    font=Font(color=font_color, bold=True),
                ),
            )

        worksheet.conditional_formatting.add(
            f"I9:I{max_row}",
            CellIsRule(
                operator="lessThan",
                formula=["0.95"],
                fill=PatternFill("solid", fgColor="F4CCCC"),
                font=Font(color="9C0006", bold=True),
            ),
        )
        worksheet.conditional_formatting.add(
            f"I9:I{max_row}",
            CellIsRule(
                operator="greaterThanOrEqual",
                formula=["0.95"],
                fill=PatternFill("solid", fgColor="E2F0D9"),
                font=Font(color="006100", bold=True),
            ),
        )

        restore_rules = (
            ('"Passed"', "E2F0D9", "006100"),
            ('"Pending"', "FFF2CC", "7F6000"),
            ('"Failed"', "F4CCCC", "9C0006"),
        )
        for formula, fill_color, font_color in restore_rules:
            worksheet.conditional_formatting.add(
                f"J9:J{max_row}",
                CellIsRule(
                    operator="equal",
                    formula=[formula],
                    fill=PatternFill("solid", fgColor=fill_color),
                    font=Font(color=font_color, bold=True),
                ),
            )

    @staticmethod
    def _apply_named_style(cell: Cell, named_style: Any) -> None:
        """Copy a named style's visual attributes onto a cell."""
        cell.font = copy(named_style.font)
        cell.fill = copy(named_style.fill)
        cell.border = copy(named_style.border)
        cell.alignment = copy(named_style.alignment)
        cell.number_format = named_style.number_format

    @staticmethod
    def _auto_fit_columns(worksheet: Worksheet) -> None:
        """Set practical column widths based on visible cell values."""
        for column_index, column_cells in enumerate(worksheet.columns, start=1):
            max_length = 0
            column_letter = get_column_letter(column_index)

            for cell in column_cells:
                if cell.value is None:
                    continue
                max_length = max(max_length, len(str(cell.value)))

            worksheet.column_dimensions[column_letter].width = min(
                max(max_length + 2, 12),
                36,
            )

    @staticmethod
    def _sample_records() -> list[BackupRecord]:
        """Return representative backup register rows."""
        return [
            BackupRecord(
                backup_job="Daily-Production",
                schedule="Daily 22:00",
                repository="Repo-Primary",
                retention="30 daily, 12 monthly",
                last_backup=date(2026, 6, 22),
                status="Success",
                successful_runs=29,
                total_runs=30,
                restore_test="Passed",
                restore_test_date=date(2026, 6, 10),
                protected_vms=42,
                owner="Backup Operations",
                notes="Production SLA compliant",
            ),
            BackupRecord(
                backup_job="Weekly-Development",
                schedule="Saturday 20:00",
                repository="Repo-Development",
                retention="8 weekly",
                last_backup=date(2026, 6, 20),
                status="Warning",
                successful_runs=7,
                total_runs=8,
                restore_test="Pending",
                restore_test_date="",
                protected_vms=24,
                owner="Infrastructure Operations",
                notes="Restore test scheduled",
            ),
            BackupRecord(
                backup_job="Critical-SQL-Logs",
                schedule="Every 4 hours",
                repository="Repo-Primary",
                retention="14 days",
                last_backup=date(2026, 6, 22),
                status="Success",
                successful_runs=180,
                total_runs=180,
                restore_test="Passed",
                restore_test_date=date(2026, 6, 5),
                protected_vms=8,
                owner="Database Services",
                notes="Transaction log backups healthy",
            ),
            BackupRecord(
                backup_job="Legacy-Test-VMs",
                schedule="Weekly Sunday 01:00",
                repository="Repo-Archive",
                retention="4 weekly",
                last_backup=date(2026, 6, 16),
                status="Failed",
                successful_runs=2,
                total_runs=4,
                restore_test="Failed",
                restore_test_date=date(2026, 6, 18),
                protected_vms=6,
                owner="Web Platform Team",
                notes="Repository access failure under investigation",
            ),
        ]


