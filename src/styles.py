"""Reusable workbook styles for Virtualization Administration Toolkit.

This module contains only styling definitions and factory functions. It does
not create worksheets, write cell values, or apply layout logic.
"""

from __future__ import annotations

from dataclasses import dataclass

from openpyxl.styles import Alignment, Border, Font, NamedStyle, PatternFill, Side


@dataclass(frozen=True)
class CorporatePalette:
    """Corporate color palette used across generated workbooks."""

    navy: str = "1F4E78"
    blue: str = "5B9BD5"
    light_blue: str = "D9EAF7"
    green: str = "70AD47"
    amber: str = "FFC000"
    red: str = "C00000"
    dark_gray: str = "404040"
    medium_gray: str = "808080"
    light_gray: str = "D9D9D9"
    white: str = "FFFFFF"


@dataclass(frozen=True)
class FontSet:
    """Reusable font definitions for workbook text hierarchy."""

    default: Font
    header: Font
    section_title: Font
    table_header: Font
    kpi_title: Font
    kpi_value: Font


@dataclass(frozen=True)
class BorderSet:
    """Reusable border definitions for workbook visual separation."""

    thin: Border
    medium_bottom: Border
    table: Border
    kpi_card: Border


@dataclass(frozen=True)
class AlignmentSet:
    """Reusable cell alignment definitions."""

    left: Alignment
    center: Alignment
    right: Alignment
    wrapped: Alignment
    kpi_center: Alignment


@dataclass(frozen=True)
class WorkbookTheme:
    """Complete workbook theme bundle for future report modules."""

    palette: CorporatePalette
    fonts: FontSet
    borders: BorderSet
    alignments: AlignmentSet
    named_styles: dict[str, NamedStyle]


def corporate_color_palette() -> CorporatePalette:
    """Return the corporate workbook color palette."""
    return CorporatePalette()


def workbook_fonts(palette: CorporatePalette | None = None) -> FontSet:
    """Return reusable workbook fonts."""
    colors = palette or corporate_color_palette()
    return FontSet(
        default=Font(name="Calibri", size=11, color=colors.dark_gray),
        header=Font(name="Calibri", size=12, bold=True, color=colors.white),
        section_title=Font(
            name="Calibri",
            size=14,
            bold=True,
            color=colors.navy,
        ),
        table_header=Font(name="Calibri", size=11, bold=True, color=colors.white),
        kpi_title=Font(name="Calibri", size=10, bold=True, color=colors.dark_gray),
        kpi_value=Font(name="Calibri", size=18, bold=True, color=colors.navy),
    )


def workbook_borders(palette: CorporatePalette | None = None) -> BorderSet:
    """Return reusable workbook border definitions."""
    colors = palette or corporate_color_palette()
    thin_gray = Side(style="thin", color=colors.light_gray)
    medium_navy = Side(style="medium", color=colors.navy)
    card_side = Side(style="thin", color=colors.medium_gray)

    return BorderSet(
        thin=Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray),
        medium_bottom=Border(bottom=medium_navy),
        table=Border(
            left=thin_gray,
            right=thin_gray,
            top=thin_gray,
            bottom=thin_gray,
        ),
        kpi_card=Border(
            left=card_side,
            right=card_side,
            top=card_side,
            bottom=card_side,
        ),
    )


def workbook_alignments() -> AlignmentSet:
    """Return reusable workbook cell alignments."""
    return AlignmentSet(
        left=Alignment(horizontal="left", vertical="center"),
        center=Alignment(horizontal="center", vertical="center"),
        right=Alignment(horizontal="right", vertical="center"),
        wrapped=Alignment(horizontal="left", vertical="top", wrap_text=True),
        kpi_center=Alignment(horizontal="center", vertical="center"),
    )


def header_style(name: str = "header") -> NamedStyle:
    """Return a named style for major worksheet headers."""
    palette = corporate_color_palette()
    fonts = workbook_fonts(palette)
    borders = workbook_borders(palette)
    alignments = workbook_alignments()

    style = NamedStyle(name=name)
    style.font = fonts.header
    style.fill = PatternFill("solid", fgColor=palette.navy)
    style.border = borders.medium_bottom
    style.alignment = alignments.center
    return style


def section_title_style(name: str = "section_title") -> NamedStyle:
    """Return a named style for section title rows."""
    palette = corporate_color_palette()
    fonts = workbook_fonts(palette)
    borders = workbook_borders(palette)
    alignments = workbook_alignments()

    style = NamedStyle(name=name)
    style.font = fonts.section_title
    style.fill = PatternFill("solid", fgColor=palette.light_blue)
    style.border = borders.medium_bottom
    style.alignment = alignments.left
    return style


def table_styles() -> dict[str, NamedStyle]:
    """Return named styles for table headers and table body cells."""
    palette = corporate_color_palette()
    fonts = workbook_fonts(palette)
    borders = workbook_borders(palette)
    alignments = workbook_alignments()

    table_header = NamedStyle(name="table_header")
    table_header.font = fonts.table_header
    table_header.fill = PatternFill("solid", fgColor=palette.blue)
    table_header.border = borders.table
    table_header.alignment = alignments.center

    table_body = NamedStyle(name="table_body")
    table_body.font = fonts.default
    table_body.fill = PatternFill("solid", fgColor=palette.white)
    table_body.border = borders.table
    table_body.alignment = alignments.left

    table_numeric = NamedStyle(name="table_numeric")
    table_numeric.font = fonts.default
    table_numeric.fill = PatternFill("solid", fgColor=palette.white)
    table_numeric.border = borders.table
    table_numeric.alignment = alignments.right

    return {
        "table_header": table_header,
        "table_body": table_body,
        "table_numeric": table_numeric,
    }


def kpi_card_style() -> dict[str, NamedStyle]:
    """Return named styles for KPI card title, value, and status cells."""
    palette = corporate_color_palette()
    fonts = workbook_fonts(palette)
    borders = workbook_borders(palette)
    alignments = workbook_alignments()

    kpi_title = NamedStyle(name="kpi_title")
    kpi_title.font = fonts.kpi_title
    kpi_title.fill = PatternFill("solid", fgColor=palette.light_blue)
    kpi_title.border = borders.kpi_card
    kpi_title.alignment = alignments.kpi_center

    kpi_value = NamedStyle(name="kpi_value")
    kpi_value.font = fonts.kpi_value
    kpi_value.fill = PatternFill("solid", fgColor=palette.white)
    kpi_value.border = borders.kpi_card
    kpi_value.alignment = alignments.kpi_center

    kpi_status = NamedStyle(name="kpi_status")
    kpi_status.font = fonts.default
    kpi_status.fill = PatternFill("solid", fgColor=palette.white)
    kpi_status.border = borders.kpi_card
    kpi_status.alignment = alignments.kpi_center

    return {
        "kpi_title": kpi_title,
        "kpi_value": kpi_value,
        "kpi_status": kpi_status,
    }


def workbook_theme() -> WorkbookTheme:
    """Return the complete reusable workbook theme."""
    palette = corporate_color_palette()
    named_styles = {
        "header": header_style(),
        "section_title": section_title_style(),
        **table_styles(),
        **kpi_card_style(),
    }

    return WorkbookTheme(
        palette=palette,
        fonts=workbook_fonts(palette),
        borders=workbook_borders(palette),
        alignments=workbook_alignments(),
        named_styles=named_styles,
    )


