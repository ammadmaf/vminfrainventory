"""VM Inventory worksheet generation.

This module creates only the ``VM Inventory`` worksheet. It does not create
formulas or any additional report worksheets.
"""

from __future__ import annotations

from copy import copy
from dataclasses import dataclass, field
from datetime import date
from typing import Any, Mapping

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from .styles import header_style, table_styles, workbook_theme


BASE_VM_INVENTORY_COLUMNS: tuple[str, ...] = (
    "VM Name",
    "Description",
    "Environment",
    "Business Unit",
    "Owner",
    "Application",
    "Operating System",
    "Host",
    "Cluster",
    "Datastore",
    "Folder",
    "Resource Pool",
    "vCPU",
    "RAM (GB)",
    "Disk (GB)",
    "IP Address",
    "VLAN",
    "Backup",
    "Backup Job",
    "Snapshot",
    "Guest Tools",
    "Power State",
    "Provision Date",
    "Last Patch Date",
    "Criticality",
    "Notes",
)

ENTERPRISE_VM_INVENTORY_COLUMNS: tuple[str, ...] = (
    "Management Server",
    "Datacenter",
    "VM MoRef ID",
    "VM UUID",
    "BIOS UUID",
    "Instance UUID",
    "Guest Hostname",
    "FQDN",
    "DNS Domain",
    "Guest OS Version",
    "Hardware Version",
    "Compatibility",
    "Guest Tools Version",
    "Guest Tools Status",
    "CPU Sockets",
    "Cores per Socket",
    "CPU Reservation (MHz)",
    "CPU Limit (MHz)",
    "CPU Shares",
    "CPU Hot Add",
    "Memory Reservation (GB)",
    "Memory Limit (GB)",
    "Memory Hot Add",
    "Memory Ballooning (MB)",
    "Guest Tools Time Sync",
    "Disk Provisioning",
    "Provisioned Disk (GB)",
    "Used Disk (GB)",
    "Free Disk (GB)",
    "Storage Policy",
    "Storage Tier",
    "Storage Path Count",
    "Storage Latency (ms)",
    "Network Adapter",
    "MAC Address",
    "Port Group",
    "Distributed Switch",
    "VLAN ID",
    "NIC Type",
    "Connected",
    "IP Allocation",
    "DNS Servers",
    "Gateway",
    "Subnet Mask",
    "HA Protection",
    "DRS Automation",
    "DRS Group",
    "Affinity Rule",
    "FT Enabled",
    "vMotion Enabled",
    "Encryption",
    "Secure Boot",
    "vTPM",
    "EVC Mode",
    "Tags",
    "Custom Attributes",
    "Cost Center",
    "Service Tier",
    "SLA",
    "Monitoring",
    "Monitoring Tool",
    "CMDB CI ID",
    "Change Request",
    "Created By",
    "Last Backup Date",
    "Last Backup Status",
    "Backup Repository",
    "Backup Policy",
    "RPO (Hours)",
    "RTO (Hours)",
    "Snapshot Count",
    "Oldest Snapshot Date",
    "Snapshot Size (GB)",
    "Replication Enabled",
    "Replication Target",
    "DR Protection Group",
    "DR Runbook",
    "Patch Group",
    "Maintenance Window",
    "Antivirus Status",
    "Compliance Status",
    "Security Zone",
    "Firewall Policy",
    "Owner Email",
    "Support Group",
    "Support Hours",
    "Decommission Date",
    "Lifecycle State",
)

VM_INVENTORY_COLUMNS: tuple[str, ...] = (
    BASE_VM_INVENTORY_COLUMNS + ENTERPRISE_VM_INVENTORY_COLUMNS
)

VM_INVENTORY_DROPDOWNS: dict[str, tuple[str, ...]] = {
    "Environment": ("Production", "Development", "Test", "Staging", "DR"),
    "Power State": ("Powered On", "Powered Off", "Suspended"),
    "Criticality": ("Critical", "High", "Medium", "Low"),
    "Backup": ("Yes", "No", "Not Required"),
    "Snapshot": ("Yes", "No"),
    "Guest Tools": ("Current", "Outdated", "Not Installed", "Not Running"),
    "CPU Hot Add": ("Enabled", "Disabled"),
    "Memory Hot Add": ("Enabled", "Disabled"),
    "Guest Tools Time Sync": ("Enabled", "Disabled"),
    "Disk Provisioning": ("Thin", "Thick Lazy Zeroed", "Thick Eager Zeroed"),
    "Connected": ("Yes", "No"),
    "IP Allocation": ("Static", "DHCP"),
    "HA Protection": ("Protected", "Not Protected"),
    "DRS Automation": ("Fully Automated", "Partially Automated", "Manual"),
    "FT Enabled": ("Yes", "No"),
    "vMotion Enabled": ("Yes", "No"),
    "Encryption": ("Encrypted", "Not Encrypted"),
    "Secure Boot": ("Enabled", "Disabled"),
    "vTPM": ("Present", "Not Present"),
    "Monitoring": ("Enabled", "Disabled"),
    "Last Backup Status": ("Success", "Warning", "Failed", "Not Configured"),
    "Replication Enabled": ("Yes", "No"),
    "Antivirus Status": ("Healthy", "Warning", "Missing"),
    "Compliance Status": ("Compliant", "Non-Compliant", "Unknown"),
    "Lifecycle State": ("Active", "Build", "Retire", "Decommissioned"),
}

VM_INVENTORY_REFERENCE_DROPDOWNS: dict[str, str] = {
    "Business Unit": "RefBusinessUnits",
    "Owner": "RefOwners",
    "Host": "RefHosts",
    "Cluster": "RefClusters",
    "Datastore": "RefDatastores",
    "Folder": "RefFolders",
    "Resource Pool": "RefResourcePools",
    "VLAN": "RefVLANs",
    "Backup Job": "RefBackupJobs",
    "Management Server": "RefManagementServers",
    "Datacenter": "RefDatacenters",
    "Port Group": "RefPortGroups",
    "Distributed Switch": "RefDistributedSwitches",
    "VLAN ID": "RefVLANs",
    "Backup Repository": "RefRepositories",
    "Security Zone": "RefSecurityZones",
    "Support Group": "RefSupportGroups",
}

NUMERIC_COLUMN_NAMES = {
    "vCPU",
    "RAM (GB)",
    "Disk (GB)",
    "CPU Sockets",
    "Cores per Socket",
    "CPU Reservation (MHz)",
    "CPU Limit (MHz)",
    "Memory Reservation (GB)",
    "Memory Limit (GB)",
    "Memory Ballooning (MB)",
    "Provisioned Disk (GB)",
    "Used Disk (GB)",
    "Free Disk (GB)",
    "Storage Path Count",
    "Storage Latency (ms)",
    "RPO (Hours)",
    "RTO (Hours)",
    "Snapshot Count",
    "Snapshot Size (GB)",
}

DATE_COLUMN_NAMES = {
    "Provision Date",
    "Last Patch Date",
    "Last Backup Date",
    "Oldest Snapshot Date",
    "Decommission Date",
}

STATUS_COLUMN_NAMES = {
    "Backup",
    "Snapshot",
    "Guest Tools",
    "Power State",
    "Criticality",
    "Connected",
    "HA Protection",
    "FT Enabled",
    "vMotion Enabled",
    "Encryption",
    "Secure Boot",
    "vTPM",
    "Monitoring",
    "Last Backup Status",
    "Replication Enabled",
    "Antivirus Status",
    "Compliance Status",
    "Lifecycle State",
}

DATA_VALIDATION_MAX_ROW = 1000
CONDITIONAL_FORMATTING_MAX_ROW = 1000

CONDITIONAL_FORMATTING_RULES: tuple[dict[str, Any], ...] = (
    {
        "column": "Power State",
        "operator": "equal",
        "formula": '"Powered On"',
        "fill": "E2F0D9",
        "font": "006100",
    },
    {
        "column": "Power State",
        "operator": "equal",
        "formula": '"Powered Off"',
        "fill": "F4CCCC",
        "font": "9C0006",
    },
    {
        "column": "Backup",
        "operator": "equal",
        "formula": '"No"',
        "fill": "F4CCCC",
        "font": "9C0006",
    },
    {
        "column": "Snapshot",
        "operator": "equal",
        "formula": '"Yes"',
        "fill": "FFF2CC",
        "font": "7F6000",
    },
    {
        "column": "Guest Tools",
        "operator": "equal",
        "formula": '"Outdated"',
        "fill": "FCE4D6",
        "font": "C65911",
    },
    {
        "column": "Criticality",
        "operator": "equal",
        "formula": '"Critical"',
        "fill": "7F0000",
        "font": "FFFFFF",
    },
    {
        "column": "RAM (GB)",
        "operator": "greaterThan",
        "formula": "64",
        "fill": "D9EAF7",
        "font": "1F4E78",
    },
    {
        "column": "Disk (GB)",
        "operator": "greaterThan",
        "formula": "1000",
        "fill": "E4DFEC",
        "font": "604A7B",
    },
)


def _enterprise_field_values(**overrides: Any) -> dict[str, Any]:
    """Return complete enterprise inventory values for a sample VM row."""
    values: dict[str, Any] = {
        "Management Server": "mgmt-primary.company.local",
        "Datacenter": "Primary Datacenter",
        "VM MoRef ID": "vm-0000",
        "VM UUID": "00000000-0000-0000-0000-000000000000",
        "BIOS UUID": "42000000-0000-0000-0000-000000000000",
        "Instance UUID": "50000000-0000-0000-0000-000000000000",
        "Guest Hostname": "hostname.company.local",
        "FQDN": "hostname.company.local",
        "DNS Domain": "company.local",
        "Guest OS Version": "Current",
        "Hardware Version": "vmx-21",
        "Compatibility": "Generic Hypervisor 2026 and later",
        "Guest Tools Version": "12.4.0",
        "Guest Tools Status": "Running",
        "CPU Sockets": 2,
        "Cores per Socket": 4,
        "CPU Reservation (MHz)": 0,
        "CPU Limit (MHz)": 0,
        "CPU Shares": "Normal",
        "CPU Hot Add": "Disabled",
        "Memory Reservation (GB)": 0,
        "Memory Limit (GB)": 0,
        "Memory Hot Add": "Disabled",
        "Memory Ballooning (MB)": 0,
        "Guest Tools Time Sync": "Disabled",
        "Disk Provisioning": "Thin",
        "Provisioned Disk (GB)": 500,
        "Used Disk (GB)": 250,
        "Free Disk (GB)": 250,
        "Storage Policy": "Gold Policy",
        "Storage Tier": "Tier 1",
        "Storage Path Count": 4,
        "Storage Latency (ms)": 3,
        "Network Adapter": "Network adapter 1",
        "MAC Address": "00:50:56:AA:00:00",
        "Port Group": "PG-Server",
        "Distributed Switch": "DSwitch-Production",
        "VLAN ID": "120",
        "NIC Type": "VMXNET3",
        "Connected": "Yes",
        "IP Allocation": "Static",
        "DNS Servers": "10.10.1.10, 10.10.1.11",
        "Gateway": "10.10.20.1",
        "Subnet Mask": "255.255.255.0",
        "HA Protection": "Protected",
        "DRS Automation": "Fully Automated",
        "DRS Group": "None",
        "Affinity Rule": "None",
        "FT Enabled": "No",
        "vMotion Enabled": "Yes",
        "Encryption": "Not Encrypted",
        "Secure Boot": "Enabled",
        "vTPM": "Not Present",
        "EVC Mode": "Intel Skylake",
        "Tags": "Managed, Production",
        "Custom Attributes": "OwnedBy=IT Operations",
        "Cost Center": "IT-000",
        "Service Tier": "Tier 2",
        "SLA": "99.5%",
        "Monitoring": "Enabled",
        "Monitoring Tool": "vROps",
        "CMDB CI ID": "CI-000000",
        "Change Request": "CHG-000000",
        "Created By": "Virtualization Automation",
        "Last Backup Date": date(2026, 6, 1),
        "Last Backup Status": "Success",
        "Backup Repository": "Repo-Primary",
        "Backup Policy": "Standard",
        "RPO (Hours)": 24,
        "RTO (Hours)": 8,
        "Snapshot Count": 0,
        "Oldest Snapshot Date": "",
        "Snapshot Size (GB)": 0,
        "Replication Enabled": "No",
        "Replication Target": "Not Configured",
        "DR Protection Group": "Not Assigned",
        "DR Runbook": "Not Assigned",
        "Patch Group": "Monthly Standard",
        "Maintenance Window": "Sunday 02:00-04:00",
        "Antivirus Status": "Healthy",
        "Compliance Status": "Compliant",
        "Security Zone": "Internal",
        "Firewall Policy": "Standard Server Policy",
        "Owner Email": "owner@company.local",
        "Support Group": "Infrastructure Operations",
        "Support Hours": "Business Hours",
        "Decommission Date": "",
        "Lifecycle State": "Active",
    }
    values.update(overrides)
    return values


@dataclass(frozen=True)
class VirtualMachineRecord:
    """Sample virtual machine inventory row."""

    vm_name: str
    description: str
    environment: str
    business_unit: str
    owner: str
    application: str
    operating_system: str
    host: str
    cluster: str
    datastore: str
    folder: str
    resource_pool: str
    vcpu: int
    ram_gb: int
    disk_gb: int
    ip_address: str
    vlan: str
    backup: str
    backup_job: str
    snapshot: str
    guest_tools: str
    power_state: str
    provision_date: date
    last_patch_date: date
    criticality: str
    notes: str
    enterprise_fields: Mapping[str, Any] = field(default_factory=dict)

    def as_row(self) -> tuple[Any, ...]:
        """Return the record in worksheet column order."""
        base_row = (
            self.vm_name,
            self.description,
            self.environment,
            self.business_unit,
            self.owner,
            self.application,
            self.operating_system,
            self.host,
            self.cluster,
            self.datastore,
            self.folder,
            self.resource_pool,
            self.vcpu,
            self.ram_gb,
            self.disk_gb,
            self.ip_address,
            self.vlan,
            self.backup,
            self.backup_job,
            self.snapshot,
            self.guest_tools,
            self.power_state,
            self.provision_date,
            self.last_patch_date,
            self.criticality,
            self.notes,
        )
        enterprise_row = tuple(
            self.enterprise_fields.get(column, "")
            for column in ENTERPRISE_VM_INVENTORY_COLUMNS
        )
        return base_row + enterprise_row


class VirtualMachineInventoryWorksheet:
    """Builds a professional VM inventory worksheet."""

    sheet_name = "VM Inventory"
    table_name = "VMInventoryTable"

    def __init__(
        self,
        records: list[VirtualMachineRecord] | None = None,
    ) -> None:
        """Initialize the worksheet builder with optional inventory records."""
        self.records = records if records is not None else self._sample_records()
        self.theme = workbook_theme()
        self.header = header_style("vm_inventory_header")
        self.table = table_styles()

    def build(self, workbook: Workbook) -> Worksheet:
        """Create and format the VM Inventory worksheet."""
        worksheet = self._get_worksheet(workbook)
        self._reset_worksheet(worksheet)
        self._configure_worksheet(worksheet)
        self._write_headers(worksheet)
        self._write_sample_data(worksheet)
        self._apply_table(worksheet)
        self._apply_dropdowns(worksheet)
        self._apply_body_formatting(worksheet)
        self._apply_conditional_formatting(worksheet)
        self._auto_fit_columns(worksheet)
        return worksheet

    def _get_worksheet(self, workbook: Workbook) -> Worksheet:
        """Return the VM inventory worksheet without adding unrelated sheets."""
        if self.sheet_name in workbook.sheetnames:
            return workbook[self.sheet_name]

        if len(workbook.worksheets) == 1 and workbook.active.max_row == 1:
            worksheet = workbook.active
            worksheet.title = self.sheet_name
            return worksheet

        return workbook.create_sheet(self.sheet_name)

    @staticmethod
    def _reset_worksheet(worksheet: Worksheet) -> None:
        """Clear worksheet content and table definitions before rebuilding."""
        for merged_range in list(worksheet.merged_cells.ranges):
            worksheet.unmerge_cells(str(merged_range))

        if worksheet.max_row:
            worksheet.delete_rows(1, worksheet.max_row)

        worksheet.tables.clear()
        worksheet.data_validations.dataValidation = []
        worksheet.conditional_formatting._cf_rules.clear()

    def _configure_worksheet(self, worksheet: Worksheet) -> None:
        """Apply sheet-level professional layout settings."""
        worksheet.sheet_view.showGridLines = False
        worksheet.freeze_panes = "A2"
        worksheet.sheet_properties.tabColor = self.theme.palette.navy

        worksheet.page_setup.orientation = "landscape"
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0

        worksheet.row_dimensions[1].height = 28

    def _write_headers(self, worksheet: Worksheet) -> None:
        """Write the exact VM inventory headers."""
        for column_index, header in enumerate(VM_INVENTORY_COLUMNS, start=1):
            cell = worksheet.cell(row=1, column=column_index, value=header)
            self._apply_named_style(cell, self.header)

    def _write_sample_data(self, worksheet: Worksheet) -> None:
        """Write sample VM inventory data without formulas."""
        for row_index, record in enumerate(self.records, start=2):
            for column_index, value in enumerate(record.as_row(), start=1):
                worksheet.cell(row=row_index, column=column_index, value=value)

    def _apply_table(self, worksheet: Worksheet) -> None:
        """Create an Excel table with filters over the inventory range."""
        last_column = get_column_letter(len(VM_INVENTORY_COLUMNS))
        table_range = f"A1:{last_column}{worksheet.max_row}"
        table = Table(displayName=self.table_name, ref=table_range)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )

        worksheet.add_table(table)

    def _apply_dropdowns(self, worksheet: Worksheet) -> None:
        """Add data validation dropdowns to managed inventory columns."""
        for column_name, options in VM_INVENTORY_DROPDOWNS.items():
            column_index = VM_INVENTORY_COLUMNS.index(column_name) + 1
            column_letter = worksheet.cell(row=1, column=column_index).column_letter
            validation = self._create_dropdown_validation(column_name, options)
            validation_range = (
                f"{column_letter}2:{column_letter}{DATA_VALIDATION_MAX_ROW}"
            )

            worksheet.add_data_validation(validation)
            validation.add(validation_range)

        for column_name, defined_name in VM_INVENTORY_REFERENCE_DROPDOWNS.items():
            column_index = VM_INVENTORY_COLUMNS.index(column_name) + 1
            column_letter = worksheet.cell(row=1, column=column_index).column_letter
            validation = self._create_reference_validation(
                column_name,
                defined_name,
            )
            validation_range = (
                f"{column_letter}2:{column_letter}{DATA_VALIDATION_MAX_ROW}"
            )
            worksheet.add_data_validation(validation)
            validation.add(validation_range)

    @staticmethod
    def _create_dropdown_validation(
        column_name: str,
        options: tuple[str, ...],
    ) -> DataValidation:
        """Create an Excel list validation for a dropdown column."""
        formula = '"' + ",".join(options) + '"'
        validation = DataValidation(
            type="list",
            formula1=formula,
            allow_blank=True,
        )
        validation.error = f"Select a valid {column_name} value from the list."
        validation.errorTitle = "Invalid Selection"
        validation.prompt = f"Choose a {column_name} value."
        validation.promptTitle = column_name
        validation.showErrorMessage = True
        validation.showInputMessage = True
        return validation

    @staticmethod
    def _create_reference_validation(
        column_name: str,
        defined_name: str,
    ) -> DataValidation:
        """Create a dropdown validation backed by Reference Data."""
        validation = DataValidation(
            type="list",
            formula1=f"={defined_name}",
            allow_blank=True,
        )
        validation.error = (
            f"Select a valid {column_name} from the Reference Data sheet."
        )
        validation.errorTitle = "Invalid Reference"
        validation.prompt = (
            f"Choose a governed {column_name} value. Add new values on "
            "Reference Data first."
        )
        validation.promptTitle = column_name
        validation.showErrorMessage = True
        validation.showInputMessage = True
        return validation

    def _apply_body_formatting(self, worksheet: Worksheet) -> None:
        """Apply corporate table styling to all sample data cells."""
        numeric_columns = self._column_indexes(NUMERIC_COLUMN_NAMES)
        date_columns = self._column_indexes(DATE_COLUMN_NAMES)
        status_columns = self._column_indexes(STATUS_COLUMN_NAMES)

        for row in worksheet.iter_rows(
            min_row=2,
            max_row=worksheet.max_row,
            min_col=1,
            max_col=len(VM_INVENTORY_COLUMNS),
        ):
            for cell in row:
                style = (
                    self.table["table_numeric"]
                    if cell.column in numeric_columns
                    else self.table["table_body"]
                )
                self._apply_named_style(cell, style)

                if cell.column in date_columns:
                    cell.number_format = "yyyy-mm-dd"

                if cell.column in status_columns:
                    self._apply_status_fill(cell)

        for row_index in range(2, worksheet.max_row + 1):
            worksheet.row_dimensions[row_index].height = 22

    @staticmethod
    def _column_indexes(column_names: set[str]) -> set[int]:
        """Return one-based column indexes for known inventory columns."""
        return {
            VM_INVENTORY_COLUMNS.index(column_name) + 1
            for column_name in column_names
            if column_name in VM_INVENTORY_COLUMNS
        }

    def _apply_conditional_formatting(self, worksheet: Worksheet) -> None:
        """Add Excel conditional formatting for key VM inventory fields."""
        for rule_config in CONDITIONAL_FORMATTING_RULES:
            column_index = VM_INVENTORY_COLUMNS.index(rule_config["column"]) + 1
            column_letter = worksheet.cell(row=1, column=column_index).column_letter
            cell_range = (
                f"{column_letter}2:"
                f"{column_letter}{CONDITIONAL_FORMATTING_MAX_ROW}"
            )

            rule = CellIsRule(
                operator=rule_config["operator"],
                formula=[rule_config["formula"]],
                fill=PatternFill("solid", fgColor=rule_config["fill"]),
                font=Font(color=rule_config["font"], bold=True),
            )
            worksheet.conditional_formatting.add(cell_range, rule)

    def _apply_status_fill(self, cell: Cell) -> None:
        """Use subtle corporate colors for status-oriented fields."""
        value = str(cell.value).lower()

        if value in {
            "yes",
            "current",
            "powered on",
            "low",
            "enabled",
            "protected",
            "success",
            "healthy",
            "compliant",
            "active",
            "present",
            "encrypted",
        }:
            cell.fill = PatternFill("solid", fgColor="E2F0D9")
        elif value in {
            "no",
            "outdated",
            "powered off",
            "critical",
            "disabled",
            "not protected",
            "failed",
            "missing",
            "non-compliant",
            "retire",
            "not configured",
            "not present",
            "not encrypted",
        }:
            cell.fill = PatternFill("solid", fgColor="F4CCCC")
        elif value in {"warning", "medium", "high", "unknown", "build"}:
            cell.fill = PatternFill("solid", fgColor="FFF2CC")

    @staticmethod
    def _apply_named_style(cell: Cell, named_style: Any) -> None:
        """Copy a named style's visual attributes onto a cell."""
        cell.font = copy(named_style.font)
        cell.fill = copy(named_style.fill)
        cell.border = copy(named_style.border)
        cell.alignment = copy(named_style.alignment)
        cell.number_format = named_style.number_format

    @staticmethod
    def _auto_fit_columns(worksheet: Worksheet) -> None:
        """Set practical column widths based on visible cell values."""
        for column_cells in worksheet.columns:
            max_length = 0
            column_letter = column_cells[0].column_letter

            for cell in column_cells:
                if cell.value is None:
                    continue
                max_length = max(max_length, len(str(cell.value)))

            worksheet.column_dimensions[column_letter].width = min(
                max(max_length + 2, 12),
                32,
            )

    @staticmethod
    def _sample_records() -> list[VirtualMachineRecord]:
        """Return representative sample inventory rows."""
        return [
            VirtualMachineRecord(
                vm_name="APP-PRD-001",
                description="Primary application server",
                environment="Production",
                business_unit="Finance",
                owner="Application Operations",
                application="ERP Portal",
                operating_system="Windows Server 2022",
                host="host-prd-01.company.local",
                cluster="Cluster-Production",
                datastore="DS-PRD-01",
                folder="/Production/Applications",
                resource_pool="RP-Production",
                vcpu=8,
                ram_gb=96,
                disk_gb=1500,
                ip_address="10.10.20.15",
                vlan="VLAN-120",
                backup="Yes",
                backup_job="Daily-Production",
                snapshot="No",
                guest_tools="Current",
                power_state="Powered On",
                provision_date=date(2025, 1, 15),
                last_patch_date=date(2026, 5, 18),
                criticality="Critical",
                notes="Tier 1 production workload",
                enterprise_fields=_enterprise_field_values(
                    **{
                        "VM MoRef ID": "vm-1015",
                        "VM UUID": "4211f8de-7d3b-4f20-9c2c-000000001015",
                        "BIOS UUID": "4211f8de-7d3b-4f20-9c2c-000000001016",
                        "Instance UUID": "5011f8de-7d3b-4f20-9c2c-000000001017",
                        "Guest Hostname": "app-prd-001.company.local",
                        "FQDN": "app-prd-001.company.local",
                        "Guest OS Version": "Windows Server 2022 Datacenter",
                        "CPU Sockets": 2,
                        "Cores per Socket": 4,
                        "CPU Hot Add": "Enabled",
                        "Memory Reservation (GB)": 32,
                        "Memory Hot Add": "Enabled",
                        "Provisioned Disk (GB)": 1500,
                        "Used Disk (GB)": 920,
                        "Free Disk (GB)": 580,
                        "Storage Policy": "Mission Critical Gold",
                        "Storage Tier": "Tier 0",
                        "Storage Latency (ms)": 2,
                        "MAC Address": "00:50:56:AA:10:15",
                        "Port Group": "PG-PRD-APP",
                        "VLAN ID": "120",
                        "Tags": "Managed, Production, Tier1",
                        "Custom Attributes": "OwnedBy=Finance; App=ERP Portal",
                        "Cost Center": "FIN-100",
                        "Service Tier": "Tier 1",
                        "SLA": "99.9%",
                        "CMDB CI ID": "CI-100015",
                        "Change Request": "CHG-260115",
                        "Last Backup Date": date(2026, 6, 22),
                        "Backup Policy": "Gold Daily",
                        "RPO (Hours)": 4,
                        "RTO (Hours)": 2,
                        "DR Protection Group": "PG-Finance-Prod",
                        "DR Runbook": "RB-FIN-ERP-001",
                        "Replication Enabled": "Yes",
                        "Replication Target": "DR Datacenter",
                        "Security Zone": "Restricted",
                        "Firewall Policy": "Finance Application Policy",
                        "Owner Email": "appops-finance@company.local",
                        "Support Group": "Finance Application Support",
                        "Support Hours": "24x7",
                    }
                ),
            ),
            VirtualMachineRecord(
                vm_name="SQL-DEV-002",
                description="Development database server",
                environment="Development",
                business_unit="Engineering",
                owner="Database Services",
                application="Analytics Sandbox",
                operating_system="Ubuntu Server 24.04",
                host="host-dev-02.company.local",
                cluster="Cluster-Development",
                datastore="DS-DEV-02",
                folder="/Development/Databases",
                resource_pool="RP-Development",
                vcpu=4,
                ram_gb=16,
                disk_gb=250,
                ip_address="10.20.30.22",
                vlan="VLAN-230",
                backup="Yes",
                backup_job="Weekly-Development",
                snapshot="Yes",
                guest_tools="Current",
                power_state="Powered On",
                provision_date=date(2025, 7, 8),
                last_patch_date=date(2026, 4, 29),
                criticality="Medium",
                notes="Snapshot pending review",
                enterprise_fields=_enterprise_field_values(
                    **{
                        "VM MoRef ID": "vm-2022",
                        "VM UUID": "4211f8de-7d3b-4f20-9c2c-000000002022",
                        "BIOS UUID": "4211f8de-7d3b-4f20-9c2c-000000002023",
                        "Instance UUID": "5011f8de-7d3b-4f20-9c2c-000000002024",
                        "Guest Hostname": "sql-dev-002.company.local",
                        "FQDN": "sql-dev-002.company.local",
                        "Guest OS Version": "Ubuntu Server 24.04 LTS",
                        "CPU Sockets": 1,
                        "Cores per Socket": 4,
                        "Provisioned Disk (GB)": 250,
                        "Used Disk (GB)": 180,
                        "Free Disk (GB)": 70,
                        "Storage Policy": "Silver Policy",
                        "Storage Tier": "Tier 2",
                        "MAC Address": "00:50:56:AA:20:22",
                        "Port Group": "PG-DEV-DB",
                        "Distributed Switch": "DSwitch-Development",
                        "VLAN ID": "230",
                        "Tags": "Managed, Development, Database",
                        "Custom Attributes": "OwnedBy=Database Services",
                        "Cost Center": "ENG-220",
                        "Service Tier": "Tier 3",
                        "SLA": "Best Effort",
                        "CMDB CI ID": "CI-200022",
                        "Change Request": "CHG-260708",
                        "Last Backup Date": date(2026, 6, 16),
                        "Backup Repository": "Repo-Development",
                        "Backup Policy": "Development Weekly",
                        "RPO (Hours)": 168,
                        "RTO (Hours)": 24,
                        "Snapshot Count": 1,
                        "Oldest Snapshot Date": date(2026, 6, 10),
                        "Snapshot Size (GB)": 18,
                        "Patch Group": "Dev Monthly",
                        "Maintenance Window": "Tuesday 21:00-23:00",
                        "Security Zone": "Development",
                        "Owner Email": "dbservices@company.local",
                        "Support Group": "Database Services",
                    }
                ),
            ),
            VirtualMachineRecord(
                vm_name="WEB-TST-003",
                description="Test web front end",
                environment="Test",
                business_unit="Digital",
                owner="Web Platform Team",
                application="Customer Portal",
                operating_system="Red Hat Enterprise Linux 9",
                host="host-tst-01.company.local",
                cluster="Cluster-Test",
                datastore="DS-TST-01",
                folder="/Test/Web",
                resource_pool="RP-Test",
                vcpu=2,
                ram_gb=8,
                disk_gb=120,
                ip_address="10.30.40.31",
                vlan="VLAN-340",
                backup="No",
                backup_job="Not Assigned",
                snapshot="No",
                guest_tools="Outdated",
                power_state="Powered Off",
                provision_date=date(2024, 11, 12),
                last_patch_date=date(2026, 3, 5),
                criticality="Low",
                notes="Scheduled for lifecycle review",
                enterprise_fields=_enterprise_field_values(
                    **{
                        "VM MoRef ID": "vm-3031",
                        "VM UUID": "4211f8de-7d3b-4f20-9c2c-000000003031",
                        "BIOS UUID": "4211f8de-7d3b-4f20-9c2c-000000003032",
                        "Instance UUID": "5011f8de-7d3b-4f20-9c2c-000000003033",
                        "Guest Hostname": "web-tst-003.company.local",
                        "FQDN": "web-tst-003.company.local",
                        "Guest OS Version": "Red Hat Enterprise Linux 9.4",
                        "Guest Tools Version": "11.3.5",
                        "Guest Tools Status": "Outdated",
                        "CPU Sockets": 1,
                        "Cores per Socket": 2,
                        "Provisioned Disk (GB)": 120,
                        "Used Disk (GB)": 62,
                        "Free Disk (GB)": 58,
                        "Storage Policy": "Bronze Policy",
                        "Storage Tier": "Tier 3",
                        "MAC Address": "00:50:56:AA:30:31",
                        "Port Group": "PG-TST-WEB",
                        "Distributed Switch": "DSwitch-Test",
                        "VLAN ID": "340",
                        "Connected": "No",
                        "Tags": "Managed, Test, Review",
                        "Custom Attributes": "Lifecycle=Review",
                        "Cost Center": "DIG-310",
                        "Service Tier": "Tier 4",
                        "SLA": "Best Effort",
                        "CMDB CI ID": "CI-300031",
                        "Change Request": "CHG-241112",
                        "Last Backup Date": "",
                        "Last Backup Status": "Not Configured",
                        "Backup Repository": "Not Configured",
                        "Backup Policy": "Not Assigned",
                        "RPO (Hours)": 0,
                        "RTO (Hours)": 0,
                        "Patch Group": "Test Monthly",
                        "Maintenance Window": "Thursday 20:00-22:00",
                        "Antivirus Status": "Warning",
                        "Compliance Status": "Unknown",
                        "Security Zone": "Test",
                        "Owner Email": "webplatform@company.local",
                        "Support Group": "Web Platform Team",
                        "Decommission Date": date(2026, 12, 31),
                        "Lifecycle State": "Retire",
                    }
                ),
            ),
        ]


