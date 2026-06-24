"""Snapshot Tracker worksheet generation.

This module creates only the ``Snapshot Tracker`` worksheet and calculates
snapshot age automatically with Excel formulas.
"""

from __future__ import annotations

from copy import copy
from dataclasses import dataclass
from datetime import date
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from .styles import header_style, table_styles, workbook_theme


SNAPSHOT_COLUMNS: tuple[str, ...] = (
    "VM Name",
    "Snapshot Name",
    "Created Date",
    "Age (Days)",
    "Created By",
    "Size (GB)",
    "Description",
    "Warning",
    "Action Required",
    "Notes",
)


@dataclass(frozen=True)
class SnapshotRecord:
    """Sample VM snapshot tracker row."""

    vm_name: str
    snapshot_name: str
    created_date: date
    created_by: str
    size_gb: float
    description: str
    notes: str

    def as_row(self, row_index: int) -> tuple[Any, ...]:
        """Return the record in worksheet order with formula columns."""
        age_formula = f"=TODAY()-C{row_index}"
        warning_formula = (
            f'=IF(D{row_index}>=90,"90+ Days",'
            f'IF(D{row_index}>=30,"30+ Days",'
            f'IF(D{row_index}>=14,"14+ Days",'
            f'IF(D{row_index}>=7,"7+ Days","Healthy"))))'
        )
        action_formula = (
            f'=IF(D{row_index}>=30,"Remove or approve retention",'
            f'IF(D{row_index}>=14,"Review owner exception",'
            f'IF(D{row_index}>=7,"Validate business need","None")))'
        )

        return (
            self.vm_name,
            self.snapshot_name,
            self.created_date,
            age_formula,
            self.created_by,
            self.size_gb,
            self.description,
            warning_formula,
            action_formula,
            self.notes,
        )


class SnapshotWorksheet:
    """Builds a professional Snapshot Tracker worksheet."""

    sheet_name = "Snapshot Tracker"
    table_name = "SnapshotTrackerTable"

    def __init__(self, records: list[SnapshotRecord] | None = None) -> None:
        """Initialize the worksheet builder with optional snapshot records."""
        self.records = records if records is not None else self._sample_records()
        self.theme = workbook_theme()
        self.header = header_style("snapshot_tracker_header")
        self.table = table_styles()

    def build(self, workbook: Workbook) -> Worksheet:
        """Create and format the Snapshot Tracker worksheet."""
        worksheet = self._get_worksheet(workbook)
        self._reset_worksheet(worksheet)
        self._configure_worksheet(worksheet)
        self._write_headers(worksheet)
        self._write_sample_data(worksheet)
        self._apply_table(worksheet)
        self._apply_dropdowns(worksheet)
        self._apply_body_formatting(worksheet)
        self._apply_conditional_formatting(worksheet)
        self._auto_fit_columns(worksheet)
        return worksheet

    def _get_worksheet(self, workbook: Workbook) -> Worksheet:
        """Return the tracker worksheet without creating unrelated sheets."""
        if self.sheet_name in workbook.sheetnames:
            return workbook[self.sheet_name]

        if len(workbook.worksheets) == 1 and workbook.active.max_row == 1:
            worksheet = workbook.active
            worksheet.title = self.sheet_name
            return worksheet

        return workbook.create_sheet(self.sheet_name)

    @staticmethod
    def _reset_worksheet(worksheet: Worksheet) -> None:
        """Clear content, tables, and conditional formatting."""
        for merged_range in list(worksheet.merged_cells.ranges):
            worksheet.unmerge_cells(str(merged_range))

        if worksheet.max_row:
            worksheet.delete_rows(1, worksheet.max_row)

        worksheet.tables.clear()
        worksheet.conditional_formatting._cf_rules.clear()

    def _configure_worksheet(self, worksheet: Worksheet) -> None:
        """Apply professional worksheet layout settings."""
        worksheet.sheet_view.showGridLines = False
        worksheet.freeze_panes = "A2"
        worksheet.sheet_properties.tabColor = self.theme.palette.amber
        worksheet.page_setup.orientation = "landscape"
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0
        worksheet.row_dimensions[1].height = 28

    def _write_headers(self, worksheet: Worksheet) -> None:
        """Write Snapshot Tracker headers."""
        for column_index, header in enumerate(SNAPSHOT_COLUMNS, start=1):
            cell = worksheet.cell(row=1, column=column_index, value=header)
            self._apply_named_style(cell, self.header)

    def _write_sample_data(self, worksheet: Worksheet) -> None:
        """Write representative snapshot records with age formulas."""
        for row_index, record in enumerate(self.records, start=2):
            for column_index, value in enumerate(record.as_row(row_index), start=1):
                worksheet.cell(row=row_index, column=column_index, value=value)

    def _apply_table(self, worksheet: Worksheet) -> None:
        """Create an Excel table with filters over snapshot data."""
        table_range = f"A1:J{worksheet.max_row}"
        table = Table(displayName=self.table_name, ref=table_range)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )

        worksheet.add_table(table)

    @staticmethod
    def _apply_dropdowns(worksheet: Worksheet) -> None:
        """Add governed dropdowns to snapshot tracker fields."""
        validations = {
            "A2:A1000": ("=RefVMNames", "Select a governed VM."),
            "E2:E1000": ("=RefOwners", "Select a governed owner."),
        }
        for cell_range, (formula, prompt) in validations.items():
            validation = DataValidation(
                type="list",
                formula1=formula,
                allow_blank=True,
            )
            validation.error = "Select a valid value from the dropdown list."
            validation.errorTitle = "Invalid Selection"
            validation.prompt = prompt
            validation.promptTitle = "Governed Dropdown"
            validation.showErrorMessage = True
            validation.showInputMessage = True
            worksheet.add_data_validation(validation)
            validation.add(cell_range)

    def _apply_body_formatting(self, worksheet: Worksheet) -> None:
        """Apply corporate styles and number formats."""
        numeric_columns = {4, 6}
        date_columns = {3}

        for row in worksheet.iter_rows(
            min_row=2,
            max_row=worksheet.max_row,
            min_col=1,
            max_col=len(SNAPSHOT_COLUMNS),
        ):
            for cell in row:
                style = (
                    self.table["table_numeric"]
                    if cell.column in numeric_columns
                    else self.table["table_body"]
                )
                self._apply_named_style(cell, style)

                if cell.column in date_columns:
                    cell.number_format = "yyyy-mm-dd"
                elif cell.column == 4:
                    cell.number_format = "0"
                elif cell.column == 6:
                    cell.number_format = "0.0"

        for row_index in range(2, worksheet.max_row + 1):
            worksheet.row_dimensions[row_index].height = 22

    @staticmethod
    def _apply_conditional_formatting(worksheet: Worksheet) -> None:
        """Highlight snapshot age and warning levels by severity."""
        max_row = max(worksheet.max_row, 2)

        rules = (
            ("greaterThanOrEqual", ["90"], "7F0000", "FFFFFF"),
            ("between", ["30", "89"], "C00000", "FFFFFF"),
            ("between", ["14", "29"], "F4B183", "7F6000"),
            ("between", ["7", "13"], "FFF2CC", "7F6000"),
        )

        for operator, formula, fill_color, font_color in rules:
            worksheet.conditional_formatting.add(
                f"D2:D{max_row}",
                CellIsRule(
                    operator=operator,
                    formula=formula,
                    fill=PatternFill("solid", fgColor=fill_color),
                    font=Font(color=font_color, bold=True),
                ),
            )

        warning_rules = (
            ('"90+ Days"', "7F0000", "FFFFFF"),
            ('"30+ Days"', "C00000", "FFFFFF"),
            ('"14+ Days"', "F4B183", "7F6000"),
            ('"7+ Days"', "FFF2CC", "7F6000"),
            ('"Healthy"', "E2F0D9", "006100"),
        )

        for formula, fill_color, font_color in warning_rules:
            worksheet.conditional_formatting.add(
                f"H2:H{max_row}",
                CellIsRule(
                    operator="equal",
                    formula=[formula],
                    fill=PatternFill("solid", fgColor=fill_color),
                    font=Font(color=font_color, bold=True),
                ),
            )

    @staticmethod
    def _apply_named_style(cell: Cell, named_style: Any) -> None:
        """Copy a named style's visual attributes onto a cell."""
        cell.font = copy(named_style.font)
        cell.fill = copy(named_style.fill)
        cell.border = copy(named_style.border)
        cell.alignment = copy(named_style.alignment)
        cell.number_format = named_style.number_format

    @staticmethod
    def _auto_fit_columns(worksheet: Worksheet) -> None:
        """Set practical column widths based on visible cell values."""
        for column_cells in worksheet.columns:
            max_length = 0
            column_letter = column_cells[0].column_letter

            for cell in column_cells:
                if cell.value is None:
                    continue
                max_length = max(max_length, len(str(cell.value)))

            worksheet.column_dimensions[column_letter].width = min(
                max(max_length + 2, 12),
                42,
            )

    @staticmethod
    def _sample_records() -> list[SnapshotRecord]:
        """Return representative snapshot tracker rows."""
        return [
            SnapshotRecord(
                vm_name="APP-PRD-001",
                snapshot_name="Pre-Patch-June",
                created_date=date(2026, 6, 20),
                created_by="Change Automation",
                size_gb=12.5,
                description="Temporary snapshot before monthly patching",
                notes="Expected to be removed after validation",
            ),
            SnapshotRecord(
                vm_name="SQL-DEV-002",
                snapshot_name="Schema-Test-Rollback",
                created_date=date(2026, 6, 12),
                created_by="Database Services",
                size_gb=34.2,
                description="Database schema testing rollback point",
                notes="Owner review required",
            ),
            SnapshotRecord(
                vm_name="WEB-TST-003",
                snapshot_name="Legacy-App-Testing",
                created_date=date(2026, 5, 18),
                created_by="Web Platform Team",
                size_gb=8.7,
                description="Application testing snapshot",
                notes="Older than 30 days",
            ),
            SnapshotRecord(
                vm_name="FILE-LEG-004",
                snapshot_name="Migration-Hold",
                created_date=date(2026, 3, 10),
                created_by="Infrastructure Operations",
                size_gb=156.4,
                description="Legacy migration hold point",
                notes="Critical cleanup required",
            ),
        ]


